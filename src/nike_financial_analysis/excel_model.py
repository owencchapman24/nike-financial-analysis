"""Build and verify the formula-driven Phase 5B Excel valuation model."""

from __future__ import annotations

import argparse
import csv
import ctypes
from datetime import date, datetime
from decimal import Decimal
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import time
from typing import Iterable, Mapping, Sequence
import xml.etree.ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import openpyxl
import xlsxwriter
from xlsxwriter.utility import xl_col_to_name

from nike_financial_analysis.analysis import (
    find_repository_root,
    hash_matches_expected,
    verify_phase2_hashes,
)
from nike_financial_analysis.forecast import verify_phase3_hashes
from nike_financial_analysis.valuation import (
    PROTECTED_PHASE4_HASHES,
    ValuationModel,
    build_valuation_model,
    calculate_terminal_bridge,
    verify_phase4_hashes,
)
from nike_financial_analysis.valuation_validation import (
    VALUATION_SOURCE_ID_PATTERN,
)


WORKBOOK_PATH = Path("model/nike_valuation_model.xlsx")
SHEET_ORDER = (
    "Cover",
    "Sources",
    "Historical",
    "Scenarios",
    "WACC",
    "DCF",
    "Sensitivity",
    "Checks",
)
SCENARIO_DISPLAY = ("Bear", "Base", "Bull")
SCENARIO_BLOCK_STARTS = {"bear": 7, "base": 32, "bull": 57}
FORECAST_YEARS = tuple(f"FY{year}" for year in range(2027, 2032))
HISTORICAL_YEARS = tuple(f"FY{year}" for year in range(2022, 2027))
FIXED_PACKAGE_TIMESTAMP = (1980, 1, 1, 0, 0, 0)
FIXED_CORE_TIMESTAMP = "2026-05-31T00:00:00Z"
MODEL_AUTHOR = "Independent portfolio project"

PROTECTED_PHASE5A_HASHES = {
    Path("config/valuation_assumptions.csv"): "7de07f936ffaf79403f1794924e83c669bb22af972a0cd8073e89373eb44f60f",
    Path("config/valuation_sources.csv"): "6d7687ce75218eb859fc7b57220d1edc7cb0e4200427e5454ab950547fc5949b",
    Path("outputs/model_exports/wacc_build.csv"): "fefe913bf1a728706a6d05e51ed21948c18b8f83eef9bd7afe3df3f129bc0076",
    Path("outputs/model_exports/valuation_cash_flows.csv"): "3f573acb52f3f494f9d51c9116ff03f9c719399f2274d0053f4291d6cf7065cd",
    Path("outputs/model_exports/valuation_summary.csv"): "8c1850ed09e95012112626d7b5c45d028fe5030b109c9a3ba0e2e87ce1b83c59",
    Path("outputs/model_exports/dcf_sensitivity.csv"): "624521449efe248a38445605be46403d5edfd4b8c27a627d25d59bff3d6f3e1c",
    Path("outputs/model_exports/valuation_validation_summary.csv"): "2e36956d3c429604b785b4c2ceb5eec2f5b067e4e4f8e8c4e8b6dd3687f5acaa",
    Path("outputs/charts/10_scenario_valuation.png"): "b3d9c953a5cb15c867c45510a739a6a4f802f4677064fa0833f39997a1b51429",
}

REQUIRED_DEFINED_NAMES = {
    "SelectedScenario",
    "ScenarioList",
    "ModelDate",
    "InformationCutoff",
    "ReferenceMarketDate",
    "ReferenceMarketPrice",
    "ClassAShares",
    "ClassBShares",
    "IncrementalDilutiveShares",
    "BasicShares",
    "DilutedProxyShares",
    "CashAndShortTermInvestments",
    "FairValueDebt",
    "CarryingDebt",
    "RedeemablePreferredStock",
    "OperatingLeaseLiabilitiesMemo",
    "RiskFreeRate",
    "EquityRiskPremium",
    "DebtDefaultSpread",
    "OperatingTaxRate",
    "TerminalGrowth",
    "CalculatedWACC",
    "CarryingDebtWACC",
    "BearValuePerShare",
    "BaseValuePerShare",
    "BullValuePerShare",
    "SelectedEnterpriseValue",
    "SelectedEquityValue",
    "SelectedValuePerShare",
    "SelectedTerminalValueShare",
    "SensitivityCenter",
    "MechanicalIntegrityStatus",
    "ApprovedSnapshotStatus",
    "WarningStatus",
    "PackageBuildControlStatus",
    "ModelStatus",
}

FORMULA_ERRORS = {
    "#REF!",
    "#VALUE!",
    "#DIV/0!",
    "#NAME?",
    "#NUM!",
    "#N/A",
    "#NULL!",
    "#SPILL!",
    "#CALC!",
}


def verify_phase5a_hashes(repository_root: Path) -> dict[str, str]:
    """Verify Phase 5A text and visual content with artifact-specific rules."""

    verified: dict[str, str] = {}
    for relative, expected in PROTECTED_PHASE5A_HASHES.items():
        path = repository_root / relative
        if not path.is_file():
            raise FileNotFoundError(f"Required Phase 5A artifact is missing: {relative}")
        if not hash_matches_expected(path, expected):
            raise ValueError(f"Protected Phase 5A artifact changed: {relative}")
        verified[relative.as_posix()] = expected
    return verified


def verify_all_protected_artifacts(repository_root: Path) -> dict[str, str]:
    """Verify the immutable Phase 2 through Phase 5A artifact set."""

    verified = verify_phase2_hashes(repository_root)
    verified.update(verify_phase3_hashes(repository_root))
    verified.update(verify_phase4_hashes(repository_root))
    verified.update(verify_phase5a_hashes(repository_root))
    return verified


def _read_csv(path: Path) -> list[dict[str, str]]:
    if not path.is_file():
        raise FileNotFoundError(f"Required workbook input is missing: {path.name}")
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def _decimal(value: str | Decimal | int) -> Decimal:
    return value if isinstance(value, Decimal) else Decimal(str(value))


def _number(value: str | Decimal | int) -> float:
    return float(_decimal(value))


def _excel_date_serial(value: datetime) -> int:
    return (value.date() - date(1899, 12, 30)).days


def _typed_value(value: str, unit: str = "") -> object:
    if value == "":
        return None
    if unit == "date" or re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        return datetime.combine(date.fromisoformat(value), datetime.min.time())
    if unit in {"decimal", "USD millions", "USD per share", "millions", "x"}:
        return float(Decimal(value))
    return value


def _forecast_index(rows: Sequence[Mapping[str, str]]) -> dict[tuple[str, str, str], Decimal | None]:
    result: dict[tuple[str, str, str], Decimal | None] = {}
    for row in rows:
        value = row.get("value", "")
        result[(row["scenario"], row["fiscal_year"], row["metric"])] = (
            None if value == "" else Decimal(value)
        )
    return result


def _format_catalog(workbook: xlsxwriter.Workbook) -> dict[str, xlsxwriter.format.Format]:
    navy = "#17365D"
    mid_blue = "#D9EAF7"
    light_blue = "#EAF2F8"
    actual_fill = "#E7E6E6"
    input_fill = "#FFF2CC"
    dark_text = "#222222"
    green = "#008000"
    blue = "#0000FF"
    gray = "#666666"
    white = "#FFFFFF"
    border = "#B4C6E7"
    base = {"font_name": "Arial", "font_size": 10, "font_color": dark_text}
    amount = '_(* #,##0.0_);_(* (#,##0.0);_(* "-"_);_(@_)'
    amount_whole = '_(* #,##0_);_(* (#,##0);_(* "-"_);_(@_)'
    percent = '_(* 0.0%_);_(* (0.0%);_(* "-"_);_(@_)'
    percent_audit = '0.000000%'
    per_share = '$0.00;($0.00);-'
    multiple = '0.00"x"'
    return {
        "title": workbook.add_format({**base, "bold": True, "font_size": 15, "bottom": 1, "bottom_color": navy}),
        "subtitle": workbook.add_format({**base, "italic": True, "font_color": gray}),
        "section": workbook.add_format({**base, "bold": True, "font_color": white, "bg_color": navy, "top": 1, "bottom": 1}),
        "header": workbook.add_format({**base, "bold": True, "font_color": white, "bg_color": navy, "align": "center", "valign": "vcenter", "border": 1, "border_color": white}),
        "subheader": workbook.add_format({**base, "bold": True, "bg_color": mid_blue, "bottom": 1, "bottom_color": border}),
        "body": workbook.add_format(base),
        "body_wrap": workbook.add_format({**base, "text_wrap": True, "valign": "top"}),
        "note": workbook.add_format({**base, "italic": True, "font_color": gray, "text_wrap": True, "valign": "top"}),
        "source": workbook.add_format({**base, "font_color": dark_text}),
        "source_wrap": workbook.add_format({**base, "font_color": dark_text, "text_wrap": True, "valign": "top"}),
        "input": workbook.add_format({**base, "font_color": blue, "bg_color": input_fill}),
        "input_percent": workbook.add_format({**base, "font_color": blue, "bg_color": input_fill, "num_format": percent}),
        "formula": workbook.add_format({**base, "font_color": "#000000"}),
        "link": workbook.add_format({**base, "font_color": green}),
        "link_percent": workbook.add_format({**base, "font_color": green, "num_format": percent}),
        "cover_output": workbook.add_format({**base, "bold": True, "font_size": 12, "font_color": dark_text}),
        "amount": workbook.add_format({**base, "num_format": amount}),
        "amount_whole": workbook.add_format({**base, "num_format": amount_whole}),
        "amount_formula": workbook.add_format({**base, "font_color": "#000000", "num_format": amount}),
        "amount_link": workbook.add_format({**base, "font_color": green, "num_format": amount}),
        "percent": workbook.add_format({**base, "num_format": percent}),
        "percent_formula": workbook.add_format({**base, "font_color": "#000000", "num_format": percent}),
        "percent_link": workbook.add_format({**base, "font_color": green, "num_format": percent}),
        "percent_audit": workbook.add_format({**base, "num_format": percent_audit}),
        "multiple": workbook.add_format({**base, "num_format": multiple}),
        "per_share": workbook.add_format({**base, "num_format": per_share}),
        "per_share_formula": workbook.add_format({**base, "font_color": "#000000", "num_format": per_share}),
        "per_share_cover": workbook.add_format({**base, "bold": True, "font_size": 12, "num_format": per_share}),
        "date": workbook.add_format({**base, "num_format": "mmm d, yyyy"}),
        "date_link": workbook.add_format({**base, "font_color": green, "num_format": "mmm d, yyyy"}),
        "actual_header": workbook.add_format({**base, "bold": True, "bg_color": actual_fill, "align": "center", "bottom": 1}),
        "estimate_header": workbook.add_format({**base, "bold": True, "bg_color": light_blue, "align": "center", "bottom": 1, "bottom_color": border}),
        "total": workbook.add_format({**base, "bold": True, "top": 1, "bottom": 2, "num_format": amount}),
        "total_per_share": workbook.add_format({**base, "bold": True, "top": 1, "bottom": 2, "num_format": per_share}),
        "pass": workbook.add_format({**base, "font_color": "#006100", "bg_color": "#C6EFCE", "bold": True, "align": "center"}),
        "warning": workbook.add_format({**base, "font_color": "#9C6500", "bg_color": "#FFEB9C", "bold": True, "align": "center"}),
        "fail": workbook.add_format({**base, "font_color": "#9C0006", "bg_color": "#FFC7CE", "bold": True, "align": "center"}),
        "nav": workbook.add_format({**base, "font_color": "#0563C1", "underline": True}),
        "legend_input": workbook.add_format({**base, "font_color": blue, "bg_color": input_fill}),
        "legend_formula": workbook.add_format({**base, "font_color": "#000000"}),
        "legend_link": workbook.add_format({**base, "font_color": green}),
    }


def _write_title(sheet: xlsxwriter.worksheet.Worksheet, title: str, subtitle: str, formats: Mapping[str, object], last_col: int) -> None:
    sheet.write(1, 1, title, formats["title"])
    sheet.set_row(1, 23)
    for col in range(2, last_col + 1):
        sheet.write_blank(1, col, None, formats["title"])
    sheet.write(2, 1, subtitle, formats["subtitle"])


def _write_formula(
    sheet: xlsxwriter.worksheet.Worksheet,
    row: int,
    col: int,
    formula: str,
    cell_format: object,
    cached: object,
) -> None:
    sheet.write_formula(row, col, formula, cell_format, cached)


def _write_sources(
    workbook: xlsxwriter.Workbook,
    sheet: xlsxwriter.worksheet.Worksheet,
    formats: Mapping[str, object],
    assumptions: list[dict[str, str]],
    sources: list[dict[str, str]],
    scenario_assumptions: list[dict[str, str]],
    model: ValuationModel,
) -> tuple[dict[str, str], dict[tuple[str, str, str], str]]:
    _write_title(sheet, "Sources and assumptions", "Embedded local inputs; no external workbook links or live data connections", formats, 13)
    sheet.hide_gridlines(2)
    sheet.set_zoom(80)
    sheet.freeze_panes(6, 0)
    sheet.set_column("A:A", 12)
    sheet.set_column("B:B", 18)
    sheet.set_column("C:C", 34)
    sheet.set_column("D:D", 17)
    sheet.set_column("E:E", 16)
    sheet.set_column("F:F", 14)
    sheet.set_column("G:G", 22)
    sheet.set_column("H:I", 22)
    sheet.set_column("J:J", 56)
    sheet.set_column("K:L", 18)
    sheet.set_column("M:M", 28)
    sheet.set_column("N:N", 35)

    input_headers = list(assumptions[0])
    input_header_labels = [name.replace("_", " ").title() for name in input_headers]
    sheet.write_row("A6", input_header_labels, formats["header"])
    input_cells: dict[str, str] = {}
    for offset, row in enumerate(assumptions, start=7):
        for col, name in enumerate(input_headers):
            value_unit = row.get("unit", "") if name == "value" else "date" if name == "observation_date" and row[name] else ""
            value = _typed_value(row[name], value_unit)
            if name == "value":
                fmt = formats["input"]
                if row.get("unit") == "decimal":
                    fmt = formats["input_percent"]
                elif row.get("unit") == "date":
                    fmt = formats["date"]
                sheet.write(offset - 1, col, value, fmt)
            elif name in {"observation_date"} and value is not None:
                sheet.write_datetime(offset - 1, col, value, formats["date"])
            else:
                sheet.write(offset - 1, col, value, formats["source_wrap"] if name == "notes" else formats["source"])
        input_cells[row["input_name"]] = f"'Sources'!$D${offset}"
    sheet.add_table("A6:J40", {"name": "tblValuationInputs", "style": "Table Style Medium 2", "columns": [{"header": label} for label in input_header_labels]})

    source_headers = list(sources[0])
    source_labels = [name.replace("_", " ").title() for name in source_headers]
    sheet.write_row("A44", source_labels, formats["header"])
    for offset, row in enumerate(sources, start=45):
        for col, name in enumerate(source_headers):
            value = _typed_value(row[name], "date" if name in {"publication_date", "observation_date", "information_cutoff", "retrieval_date"} and row[name] else "")
            if isinstance(value, datetime):
                sheet.write_datetime(offset - 1, col, value, formats["date"])
            else:
                fmt = formats["source_wrap"] if name in {"title", "url_or_accession", "fields_supported", "notes"} else formats["source"]
                sheet.write(offset - 1, col, value, fmt)
        sheet.set_row(offset - 1, 31)
    sheet.add_table("A44:L55", {"name": "tblValuationSources", "style": "Table Style Medium 2", "columns": [{"header": label} for label in source_labels]})

    scenario_headers = list(scenario_assumptions[0])
    scenario_labels = [name.replace("_", " ").title() for name in scenario_headers]
    sheet.write_row("A59", scenario_labels, formats["header"])
    scenario_cells: dict[tuple[str, str, str], str] = {}
    for offset, row in enumerate(scenario_assumptions, start=60):
        for col, name in enumerate(scenario_headers):
            value = _typed_value(row[name], row.get("unit", "") if name == "value" else "")
            if name == "value":
                sheet.write(offset - 1, col, value, formats["input_percent"])
            else:
                fmt = formats["source_wrap"] if name in {"historical_anchor", "rationale", "notes"} else formats["source"]
                sheet.write(offset - 1, col, value, fmt)
        scenario_cells[(row["scenario"], row["fiscal_year"], row["driver"])] = f"'Sources'!$D${offset}"
    sheet.add_table("A59:K164", {"name": "tblScenarioAssumptions", "style": "Table Style Medium 2", "columns": [{"header": label} for label in scenario_labels]})

    sheet.write("M6", "Scenario selector values", formats["subheader"])
    for row, scenario in enumerate(SCENARIO_DISPLAY, start=7):
        sheet.write(row - 1, 13, scenario, formats["body"])
    sheet.write("M12", "Workbook metadata", formats["subheader"])
    metadata = (
        (
            "DCF valuation date",
            datetime.combine(date.fromisoformat(model.model_date), datetime.min.time()),
            formats["date"],
        ),
        ("Information cutoff", datetime(2026, 9, 7), formats["date"]),
        ("Reference market date", datetime(2026, 9, 4), formats["date"]),
        ("Valuation inputs", "config/valuation_assumptions.csv", formats["source"]),
        ("Scenario assumptions", "config/scenario_assumptions.csv", formats["source"]),
        ("Python valuation results", "outputs/model_exports/valuation_summary.csv", formats["source"]),
    )
    for offset, (label, value, fmt) in enumerate(metadata, start=13):
        sheet.write(offset - 1, 12, label, formats["body"])
        if isinstance(value, datetime):
            sheet.write_datetime(offset - 1, 13, value, fmt)
        else:
            sheet.write(offset - 1, 13, value, fmt)
    sheet.write("M21", "Calculated input combinations", formats["subheader"])
    sheet.write("M22", "Basic shares", formats["body"])
    _write_formula(sheet, 21, 13, "=ClassAShares+ClassBShares", formats["formula"], _number(model.basic_shares))
    sheet.write("M23", "Diluted proxy shares", formats["body"])
    _write_formula(sheet, 22, 13, "=BasicShares+IncrementalDilutiveShares", formats["formula"], _number(model.diluted_proxy_shares))
    sheet.write("M24", "Cash and short-term investments", formats["body"])
    _write_formula(sheet, 23, 13, "='Sources'!D13+'Sources'!D14", formats["amount_formula"], _number(model.cash_and_short_term_investments))

    direct_names = {
        "ModelDate": "model_date",
        "ReferenceMarketDate": "reference_market_date",
        "ReferenceMarketPrice": "reference_market_price",
        "ClassAShares": "class_a_shares_outstanding",
        "ClassBShares": "class_b_shares_outstanding",
        "IncrementalDilutiveShares": "incremental_dilutive_share_proxy",
        "FairValueDebt": "fair_value_debt",
        "CarryingDebt": "carrying_interest_bearing_debt",
        "RedeemablePreferredStock": "redeemable_preferred_stock",
        "OperatingLeaseLiabilitiesMemo": "operating_lease_liabilities",
        "RiskFreeRate": "risk_free_rate",
        "EquityRiskPremium": "equity_risk_premium",
        "DebtDefaultSpread": "debt_default_spread",
        "OperatingTaxRate": "operating_tax_rate",
        "TerminalGrowth": "terminal_growth_rate",
    }
    for defined_name, input_name in direct_names.items():
        workbook.define_name(defined_name, f"={input_cells[input_name]}")
    workbook.define_name("InformationCutoff", "='Sources'!$N$14")
    workbook.define_name("ScenarioList", "='Sources'!$N$7:$N$9")
    workbook.define_name("BasicShares", "='Sources'!$N$22")
    workbook.define_name("DilutedProxyShares", "='Sources'!$N$23")
    workbook.define_name("CashAndShortTermInvestments", "='Sources'!$N$24")
    return input_cells, scenario_cells


def _write_historical(
    sheet: xlsxwriter.worksheet.Worksheet,
    formats: Mapping[str, object],
    forecast: Mapping[tuple[str, str, str], Decimal | None],
    assumptions_by_name: Mapping[str, dict[str, str]],
    model: ValuationModel,
) -> dict[str, int]:
    _write_title(sheet, "Historical DCF inputs", "FY2022A-FY2026A; USD millions except per-share, share-count, and percentage lines", formats, 9)
    sheet.hide_gridlines(2)
    sheet.set_zoom(90)
    sheet.freeze_panes(7, 3)
    sheet.set_column("A:A", 2)
    sheet.set_column("B:B", 38)
    sheet.set_column("C:C", 16)
    sheet.set_column("D:H", 14)
    sheet.set_column("I:I", 56)
    sheet.write_row("B7", ["Metric", "Unit", *[f"{year}A" for year in HISTORICAL_YEARS], "Source / convention"], formats["header"])
    rows = (
        ("revenue", "Revenue", "USD millions"),
        ("derived_operating_income", "Derived operating income", "USD millions"),
        ("operating_margin", "Derived operating margin", "decimal"),
        ("normalized_tax_rate", "Historical effective operating tax rate", "decimal"),
        ("depreciation_and_amortization", "Depreciation and amortization", "USD millions"),
        ("capital_expenditures", "Capital expenditures", "USD millions"),
        ("operating_nwc_proxy", "Operating NWC proxy", "USD millions"),
        ("change_in_operating_nwc", "Change in operating NWC", "USD millions"),
        ("fcff", "Historical FCFF", "USD millions"),
    )
    row_map: dict[str, int] = {}
    for row_number, (metric, label, unit) in enumerate(rows, start=8):
        row_map[metric] = row_number
        sheet.write(row_number - 1, 1, label, formats["body"])
        sheet.write(row_number - 1, 2, unit, formats["body"])
        for col, year in enumerate(HISTORICAL_YEARS, start=3):
            value = forecast.get(("actual", year, metric))
            if value is None:
                sheet.write_blank(row_number - 1, col, None, formats["source"])
            else:
                fmt = formats["percent"] if unit == "decimal" else formats["amount"]
                sheet.write_number(row_number - 1, col, _number(value), fmt)
        note = "Committed Phase 4 actual bridge."
        if metric == "derived_operating_income":
            note = "Project-derived consolidated subtotal: gross profit less total SG&A; not Nike-reported or segment EBIT."
        elif metric in {"change_in_operating_nwc", "fcff"}:
            note = "FY2022 is not applicable because FY2021 opening operating NWC is unavailable."
        elif metric == "capital_expenditures":
            note = "Positive amount represents investment spending."
        sheet.write(row_number - 1, 8, note, formats["note"])

    sheet.write("B20", "FY2026 valuation bridge and share inputs", formats["section"])
    for col in range(2, 9):
        sheet.write_blank(19, col, None, formats["section"])
    bridge = (
        ("Cash and cash equivalents", "USD millions", "cash_and_cash_equivalents", assumptions_by_name["cash_and_cash_equivalents"]["source_ids"]),
        ("Short-term investments", "USD millions", "short_term_investments", assumptions_by_name["short_term_investments"]["source_ids"]),
        ("Carrying interest-bearing debt", "USD millions", "carrying_interest_bearing_debt", assumptions_by_name["carrying_interest_bearing_debt"]["source_ids"]),
        ("Operating leases (memorandum)", "USD millions", "operating_lease_liabilities", assumptions_by_name["operating_lease_liabilities"]["source_ids"]),
        ("Class A shares outstanding", "millions", "class_a_shares_outstanding", assumptions_by_name["class_a_shares_outstanding"]["source_ids"]),
        ("Class B shares outstanding", "millions", "class_b_shares_outstanding", assumptions_by_name["class_b_shares_outstanding"]["source_ids"]),
        ("Incremental dilutive-share proxy", "millions", "incremental_dilutive_share_proxy", assumptions_by_name["incremental_dilutive_share_proxy"]["source_ids"]),
        ("Basic shares", "millions", None, "VA04;VA05"),
        ("Diluted proxy shares", "millions", None, "VA04;VA05;VA06"),
    )
    for row_number, (label, unit, input_name, source_ids) in enumerate(bridge, start=22):
        sheet.write(row_number - 1, 1, label, formats["body"])
        sheet.write(row_number - 1, 2, unit, formats["body"])
        if input_name:
            formula = "=" + {
                "cash_and_cash_equivalents": "'Sources'!D13",
                "short_term_investments": "'Sources'!D14",
                "carrying_interest_bearing_debt": "CarryingDebt",
                "operating_lease_liabilities": "OperatingLeaseLiabilitiesMemo",
                "class_a_shares_outstanding": "ClassAShares",
                "class_b_shares_outstanding": "ClassBShares",
                "incremental_dilutive_share_proxy": "IncrementalDilutiveShares",
            }[input_name]
            cached = assumptions_by_name[input_name]["value"]
        elif label == "Basic shares":
            formula, cached = "=BasicShares", model.basic_shares
        else:
            formula, cached = "=DilutedProxyShares", model.diluted_proxy_shares
        _write_formula(sheet, row_number - 1, 7, formula, formats["amount_link"], _number(cached))
        sheet.write(row_number - 1, 8, source_ids, formats["source"])
    return row_map


SCENARIO_METRICS = (
    ("revenue_growth", "Revenue growth", "decimal", True),
    ("revenue", "Revenue", "USD millions", False),
    ("gross_margin", "Gross margin", "decimal", True),
    ("gross_profit", "Gross profit", "USD millions", False),
    ("sga_percent_revenue", "SG&A / revenue", "decimal", True),
    ("total_selling_and_administrative_expense", "Total SG&A", "USD millions", False),
    ("derived_operating_income", "Derived operating income", "USD millions", False),
    ("operating_margin", "Derived operating margin", "decimal", False),
    ("normalized_tax_rate", "Operating tax rate", "decimal", True),
    ("normalized_operating_tax_expense", "Normalized operating tax expense", "USD millions", False),
    ("nopat", "NOPAT", "USD millions", False),
    ("da_percent_revenue", "D&A / revenue", "decimal", True),
    ("depreciation_and_amortization", "Depreciation and amortization", "USD millions", False),
    ("capex_percent_revenue", "Capex / revenue", "decimal", True),
    ("capital_expenditures", "Capital expenditures", "USD millions", False),
    ("operating_nwc_percent_revenue", "Operating NWC / revenue", "decimal", True),
    ("operating_nwc_proxy", "Operating NWC proxy", "USD millions", False),
    ("change_in_operating_nwc", "Change in operating NWC", "USD millions", False),
    ("fcff", "Free cash flow to the firm", "USD millions", False),
)


def _scenario_formula(metric: str, col_letter: str, prior_col: str, row_map: Mapping[str, int]) -> str:
    cell = lambda name, col=col_letter: f"{col}{row_map[name]}"
    formulas = {
        "revenue": f"={cell('revenue', prior_col)}*(1+{cell('revenue_growth')})",
        "gross_profit": f"={cell('revenue')}*{cell('gross_margin')}",
        "total_selling_and_administrative_expense": f"={cell('revenue')}*{cell('sga_percent_revenue')}",
        "derived_operating_income": f"={cell('gross_profit')}-{cell('total_selling_and_administrative_expense')}",
        "operating_margin": f"={cell('derived_operating_income')}/{cell('revenue')}",
        "normalized_operating_tax_expense": f"=MAX({cell('derived_operating_income')},0)*{cell('normalized_tax_rate')}",
        "nopat": f"={cell('derived_operating_income')}-{cell('normalized_operating_tax_expense')}",
        "depreciation_and_amortization": f"={cell('revenue')}*{cell('da_percent_revenue')}",
        "capital_expenditures": f"={cell('revenue')}*{cell('capex_percent_revenue')}",
        "operating_nwc_proxy": f"={cell('revenue')}*{cell('operating_nwc_percent_revenue')}",
        "change_in_operating_nwc": f"={cell('operating_nwc_proxy')}-{cell('operating_nwc_proxy', prior_col)}",
        "fcff": f"={cell('nopat')}+{cell('depreciation_and_amortization')}-{cell('capital_expenditures')}-{cell('change_in_operating_nwc')}",
    }
    return formulas[metric]


def _write_scenarios(
    sheet: xlsxwriter.worksheet.Worksheet,
    formats: Mapping[str, object],
    forecast: Mapping[tuple[str, str, str], Decimal | None],
    scenario_cells: Mapping[tuple[str, str, str], str],
) -> dict[str, dict[str, int]]:
    _write_title(sheet, "Operating scenarios", "FY2026A anchor and approved FY2027E-FY2031E analyst scenarios; USD millions unless noted", formats, 9)
    sheet.hide_gridlines(2)
    sheet.set_zoom(85)
    sheet.freeze_panes(8, 3)
    sheet.set_column("A:A", 2)
    sheet.set_column("B:B", 36)
    sheet.set_column("C:C", 14)
    sheet.set_column("D:I", 14)
    sheet.set_column("J:J", 52)
    row_maps: dict[str, dict[str, int]] = {}
    for scenario, start in SCENARIO_BLOCK_STARTS.items():
        sheet.write(start - 1, 1, f"{scenario.title()} scenario", formats["section"])
        for col in range(2, 10):
            sheet.write_blank(start - 1, col, None, formats["section"])
        sheet.write_row(start, 1, ["Metric", "Unit", "FY2026A", *[f"{year}E" for year in FORECAST_YEARS], "Source / formula"], formats["header"])
        row_map = {
            metric: start + offset
            for offset, (metric, _, _, _) in enumerate(SCENARIO_METRICS, start=2)
        }
        for offset, (metric, label, unit, driver) in enumerate(SCENARIO_METRICS, start=2):
            excel_row = start + offset
            sheet.write(excel_row - 1, 1, label, formats["body"])
            sheet.write(excel_row - 1, 2, unit, formats["body"])
            actual = forecast[("actual", "FY2026", metric)]
            actual_fmt = formats["percent"] if unit == "decimal" else formats["amount"]
            if actual is not None:
                sheet.write_number(excel_row - 1, 3, _number(actual), actual_fmt)
            for year_index, year in enumerate(FORECAST_YEARS, start=4):
                expected = forecast[(scenario, year, metric)]
                col_letter = xl_col_to_name(year_index)
                prior_col = xl_col_to_name(year_index - 1)
                if driver:
                    source = scenario_cells[(scenario, year, metric)]
                    formula = f"={source}"
                    fmt = formats["percent_link"]
                else:
                    formula = _scenario_formula(metric, col_letter, prior_col, row_map)
                    fmt = formats["amount_formula"] if unit == "USD millions" else formats["percent_formula"]
                _write_formula(sheet, excel_row - 1, year_index, formula, fmt, _number(expected))
            if metric == "derived_operating_income":
                note = "Gross profit less total SG&A; not Nike-reported or segment EBIT."
            elif driver:
                note = "Approved scenario driver linked from Sources."
            elif metric == "fcff":
                note = "NOPAT + D&A - capex - change in operating NWC."
            else:
                note = "Native workbook formula."
            sheet.write(excel_row - 1, 9, note, formats["note"])
        row_maps[scenario] = row_map
    return row_maps


def _write_wacc(
    workbook: xlsxwriter.Workbook,
    sheet: xlsxwriter.worksheet.Worksheet,
    formats: Mapping[str, object],
    model: ValuationModel,
) -> dict[str, int]:
    _write_title(sheet, "Weighted average cost of capital", "Primary build uses fair-value debt; carrying-debt build is a cross-check", formats, 9)
    sheet.hide_gridlines(2)
    sheet.set_zoom(90)
    sheet.freeze_panes(8, 1)
    sheet.set_column("A:A", 2)
    sheet.set_column("B:B", 32)
    sheet.set_column("C:C", 17)
    sheet.set_column("D:D", 25)
    sheet.set_column("E:E", 18)
    sheet.set_column("F:F", 3)
    sheet.set_column("G:G", 32)
    sheet.set_column("H:H", 17)
    sheet.set_column("I:I", 25)
    sheet.set_column("J:J", 18)
    headers = ("Metric", "Unit", "Source / formula", "Value")
    sheet.write_row("B7", ["Primary fair-value-debt WACC"], formats["section"])
    sheet.write_row("G7", ["Carrying-debt WACC cross-check"], formats["section"])
    for col in range(2, 5):
        sheet.write_blank(6, col, None, formats["section"])
    for col in range(7, 10):
        sheet.write_blank(6, col, None, formats["section"])
    sheet.write_row("B8", headers, formats["header"])
    sheet.write_row("G8", headers, formats["header"])

    metrics = (
        ("market_price", "Reference market price", "USD per share"),
        ("class_a", "Class A shares", "millions"),
        ("class_b", "Class B shares", "millions"),
        ("basic_shares", "Basic shares", "millions"),
        ("market_equity", "Market equity", "USD millions"),
        ("footwear_revenue", "Footwear revenue", "USD millions"),
        ("footwear_beta", "Footwear unlevered beta", "x"),
        ("apparel_revenue", "Apparel revenue", "USD millions"),
        ("apparel_beta", "Apparel unlevered beta", "x"),
        ("equipment_revenue", "Equipment revenue", "USD millions"),
        ("equipment_beta", "Equipment unlevered beta", "x"),
        ("classified_revenue", "Classified product revenue", "USD millions"),
        ("other_revenue", "Other revenue", "USD millions"),
        ("total_revenue", "Total revenue", "USD millions"),
        ("unlevered_beta", "Product-weighted unlevered beta", "x"),
        ("debt", "Debt used", "USD millions"),
        ("de_ratio", "Debt / equity", "decimal"),
        ("tax_rate", "Operating tax rate", "decimal"),
        ("levered_beta", "Relevered beta", "x"),
        ("risk_free", "Risk-free rate", "decimal"),
        ("erp", "Equity-risk premium", "decimal"),
        ("cost_equity", "Cost of equity", "decimal"),
        ("spread", "Debt default spread", "decimal"),
        ("pretax_debt", "Pretax cost of debt", "decimal"),
        ("after_tax_debt", "After-tax cost of debt", "decimal"),
        ("equity_weight", "Equity weight", "decimal"),
        ("debt_weight", "Debt weight", "decimal"),
        ("wacc", "Calculated WACC", "decimal"),
    )
    row_map = {metric: 9 + index for index, (metric, _, _) in enumerate(metrics)}

    for start_col, build, debt_name in ((1, model.wacc_primary, "FairValueDebt"), (6, model.wacc_carrying_cross_check, "CarryingDebt")):
        value_col = xl_col_to_name(start_col + 3)
        def cell(metric: str) -> str:
            return f"{value_col}{row_map[metric]}"
        formulas = {
            "market_price": "=ReferenceMarketPrice",
            "class_a": "=ClassAShares",
            "class_b": "=ClassBShares",
            "basic_shares": f"={cell('class_a')}+{cell('class_b')}",
            "market_equity": f"={cell('market_price')}*{cell('basic_shares')}",
            "footwear_revenue": "='Sources'!D23",
            "footwear_beta": "='Sources'!D27",
            "apparel_revenue": "='Sources'!D24",
            "apparel_beta": "='Sources'!D28",
            "equipment_revenue": "='Sources'!D25",
            "equipment_beta": "='Sources'!D29",
            "classified_revenue": f"={cell('footwear_revenue')}+{cell('apparel_revenue')}+{cell('equipment_revenue')}",
            "other_revenue": "='Sources'!D26",
            "total_revenue": f"={cell('classified_revenue')}+{cell('other_revenue')}",
            "unlevered_beta": f"=({cell('footwear_revenue')}*{cell('footwear_beta')}+{cell('apparel_revenue')}*{cell('apparel_beta')}+{cell('equipment_revenue')}*{cell('equipment_beta')})/{cell('classified_revenue')}",
            "debt": f"={debt_name}",
            "de_ratio": f"={cell('debt')}/{cell('market_equity')}",
            "tax_rate": "=OperatingTaxRate",
            "levered_beta": f"={cell('unlevered_beta')}*(1+(1-{cell('tax_rate')})*{cell('debt')}/{cell('market_equity')})",
            "risk_free": "=RiskFreeRate",
            "erp": "=EquityRiskPremium",
            "cost_equity": f"={cell('risk_free')}+{cell('levered_beta')}*{cell('erp')}",
            "spread": "=DebtDefaultSpread",
            "pretax_debt": f"={cell('risk_free')}+{cell('spread')}",
            "after_tax_debt": f"={cell('pretax_debt')}*(1-{cell('tax_rate')})",
            "equity_weight": f"={cell('market_equity')}/({cell('market_equity')}+{cell('debt')})",
            "debt_weight": f"={cell('debt')}/({cell('market_equity')}+{cell('debt')})",
            "wacc": f"={cell('equity_weight')}*{cell('cost_equity')}+{cell('debt_weight')}*{cell('after_tax_debt')}",
        }
        values = {
            "market_price": build.market_price,
            "class_a": Decimal("281.387752"),
            "class_b": Decimal("1202.110951"),
            "basic_shares": build.basic_shares,
            "market_equity": build.market_equity,
            "footwear_revenue": Decimal("30538"),
            "footwear_beta": Decimal("1.00"),
            "apparel_revenue": Decimal("13497"),
            "apparel_beta": Decimal("0.79"),
            "equipment_revenue": Decimal("2220"),
            "equipment_beta": Decimal("0.74"),
            "classified_revenue": build.classified_revenue,
            "other_revenue": build.other_revenue,
            "total_revenue": build.total_revenue,
            "unlevered_beta": build.unlevered_beta,
            "debt": build.debt,
            "de_ratio": build.debt / build.market_equity,
            "tax_rate": build.tax_rate,
            "levered_beta": build.levered_beta,
            "risk_free": build.risk_free_rate,
            "erp": build.equity_risk_premium,
            "cost_equity": build.cost_of_equity,
            "spread": build.default_spread,
            "pretax_debt": build.pretax_cost_of_debt,
            "after_tax_debt": build.after_tax_cost_of_debt,
            "equity_weight": build.equity_weight,
            "debt_weight": build.debt_weight,
            "wacc": build.wacc,
        }
        for metric, label, unit in metrics:
            excel_row = row_map[metric]
            sheet.write(excel_row - 1, start_col, label, formats["body"])
            sheet.write(excel_row - 1, start_col + 1, unit, formats["body"])
            sheet.write(excel_row - 1, start_col + 2, formulas[metric].lstrip("="), formats["note"])
            if unit == "decimal":
                fmt = formats["percent_audit"] if metric in {"wacc", "cost_equity"} else formats["percent_formula"]
            elif unit == "USD per share":
                fmt = formats["per_share_formula"]
            elif unit == "x":
                fmt = formats["multiple"]
            else:
                fmt = formats["amount_formula"]
            if formulas[metric].startswith("='Sources'") or formulas[metric] in {"=ReferenceMarketPrice", "=ClassAShares", "=ClassBShares", "=FairValueDebt", "=CarryingDebt", "=OperatingTaxRate", "=RiskFreeRate", "=EquityRiskPremium", "=DebtDefaultSpread"}:
                if unit == "decimal":
                    fmt = formats["percent_link"]
                elif unit == "USD per share":
                    fmt = formats["per_share"]
                else:
                    fmt = formats["amount_link"]
            _write_formula(sheet, excel_row - 1, start_col + 3, formulas[metric], fmt, _number(values[metric]))
        sheet.set_row(row_map["wacc"] - 1, 19)
    workbook.define_name("CalculatedWACC", f"='WACC'!$E${row_map['wacc']}")
    workbook.define_name("CarryingDebtWACC", f"='WACC'!$J${row_map['wacc']}")
    return row_map


def _selected_formula(scenario_rows: Mapping[str, Mapping[str, int]], metric: str, column: str) -> str:
    return (
        f"=CHOOSE(MATCH(SelectedScenario,ScenarioList,0),"
        f"'Scenarios'!{column}{scenario_rows['bear'][metric]},"
        f"'Scenarios'!{column}{scenario_rows['base'][metric]},"
        f"'Scenarios'!{column}{scenario_rows['bull'][metric]})"
    )


def _write_dcf(
    workbook: xlsxwriter.Workbook,
    sheet: xlsxwriter.worksheet.Worksheet,
    formats: Mapping[str, object],
    model: ValuationModel,
    forecast: Mapping[tuple[str, str, str], Decimal | None],
    scenario_rows: Mapping[str, Mapping[str, int]],
) -> dict[str, object]:
    _write_title(
        sheet,
        "Discounted cash flow",
        "Full-year FY2027-FY2031 FCFF discounted from Sep. 4, 2026; May 31, 2026 equity-bridge balances are not rolled forward",
        formats,
        13,
    )
    sheet.hide_gridlines(2)
    sheet.set_zoom(85)
    sheet.freeze_panes(10, 3)
    sheet.set_column("A:A", 2)
    sheet.set_column("B:B", 36)
    sheet.set_column("C:C", 16)
    sheet.set_column("D:I", 14)
    sheet.set_column("J:J", 3)
    sheet.set_column("K:K", 34)
    sheet.set_column("L:N", 17)

    sheet.write("B6", "Selected scenario", formats["subheader"])
    _write_formula(sheet, 5, 3, "=SelectedScenario", formats["link"], "Base")
    sheet.write("B9", "Explicit forecast", formats["section"])
    for col in range(2, 9):
        sheet.write_blank(8, col, None, formats["section"])
    sheet.write_row("B10", ["Metric", "Unit", "Model date", *[f"{year}E" for year in FORECAST_YEARS]], formats["header"])
    model_date = date.fromisoformat(model.model_date)
    date_values = [model_date, *[date(year, 5, 31) for year in range(2027, 2032)]]
    sheet.write("B12", "Cash-flow date", formats["body"])
    sheet.write("C12", "date", formats["body"])
    for col, value in enumerate(date_values, start=3):
        sheet.write_datetime(11, col, datetime.combine(value, datetime.min.time()), formats["date"])
    sheet.write("B13", "FCFF", formats["body"])
    sheet.write("C13", "USD millions", formats["body"])
    sheet.write_number("D13", 0, formats["amount"])
    base_valuation = next(item for item in model.scenario_valuations if item.scenario == "base")
    for index, col in enumerate(range(4, 9)):
        col_letter = xl_col_to_name(col)
        _write_formula(sheet, 12, col, _selected_formula(scenario_rows, "fcff", col_letter), formats["amount_link"], _number(base_valuation.explicit_cash_flows[index].fcff))
    sheet.write("B14", "Discount exponent", formats["body"])
    sheet.write("C14", "x", formats["body"])
    for col in range(3, 9):
        letter = xl_col_to_name(col)
        expected = Decimal((date_values[col - 3] - model_date).days) / Decimal("365")
        _write_formula(sheet, 13, col, f"=({letter}12-ModelDate)/365", formats["formula"], _number(expected))
    sheet.write("B15", "Discount factor", formats["body"])
    sheet.write("C15", "x", formats["body"])
    for col in range(3, 9):
        letter = xl_col_to_name(col)
        expected = Decimal("1") if col == 3 else base_valuation.explicit_cash_flows[col - 4].discount_factor
        _write_formula(sheet, 14, col, f"=1/(1+CalculatedWACC)^{letter}14", formats["formula"], _number(expected))
    sheet.write("B16", "Present value of FCFF", formats["body"])
    sheet.write("C16", "USD millions", formats["body"])
    for col in range(3, 9):
        letter = xl_col_to_name(col)
        expected = Decimal("0") if col == 3 else base_valuation.explicit_cash_flows[col - 4].present_value
        _write_formula(sheet, 15, col, f"={letter}13*{letter}15", formats["amount_formula"], _number(expected))
    sheet.write("B18", "Total explicit-period PV", formats["body"])
    _write_formula(sheet, 17, 8, "=SUM(E16:I16)", formats["total"], _number(base_valuation.pv_explicit_fcff))

    sheet.write("B21", "FY2032 terminal transition", formats["section"])
    for col in range(2, 5):
        sheet.write_blank(20, col, None, formats["section"])
    sheet.write_row("B22", ["Metric", "Unit", "FY2031E", "FY2032 terminal"], formats["header"])
    terminal = base_valuation.terminal
    terminal_rows = (
        ("revenue_growth", "Revenue growth", "decimal", terminal.revenue_growth),
        ("revenue", "Revenue", "USD millions", terminal.revenue),
        ("gross_margin", "Gross margin", "decimal", terminal.gross_margin),
        ("gross_profit", "Gross profit", "USD millions", terminal.gross_profit),
        ("sga_percent_revenue", "SG&A / revenue", "decimal", terminal.sga_percent_revenue),
        ("total_selling_and_administrative_expense", "Total SG&A", "USD millions", terminal.total_selling_and_administrative_expense),
        ("derived_operating_income", "Derived operating income", "USD millions", terminal.derived_operating_income),
        ("operating_margin", "Derived operating margin", "decimal", terminal.operating_margin),
        ("normalized_tax_rate", "Operating tax rate", "decimal", terminal.normalized_tax_rate),
        ("normalized_operating_tax_expense", "Normalized operating tax expense", "USD millions", terminal.derived_operating_income - terminal.nopat),
        ("nopat", "NOPAT", "USD millions", terminal.nopat),
        ("da_percent_revenue", "D&A / revenue", "decimal", terminal.da_percent_revenue),
        ("depreciation_and_amortization", "Depreciation and amortization", "USD millions", terminal.depreciation_and_amortization),
        ("capex_percent_revenue", "Capex / revenue", "decimal", terminal.capex_percent_revenue),
        ("capital_expenditures", "Capital expenditures", "USD millions", terminal.capital_expenditures),
        ("operating_nwc_percent_revenue", "Operating NWC / revenue", "decimal", terminal.operating_nwc_percent_revenue),
        ("operating_nwc_proxy", "Operating NWC proxy", "USD millions", terminal.operating_nwc_proxy),
        ("change_in_operating_nwc", "Change in operating NWC", "USD millions", terminal.change_in_operating_nwc),
        ("fcff", "Terminal-year FCFF", "USD millions", terminal.fcff),
        ("ebitda_proxy", "Terminal EBITDA proxy", "USD millions", terminal.ebitda_proxy),
    )
    terminal_row_map = {
        metric: offset
        for offset, (metric, _, _, _) in enumerate(terminal_rows, start=23)
    }
    for offset, (metric, label, unit, cached_terminal) in enumerate(terminal_rows, start=23):
        sheet.write(offset - 1, 1, label, formats["body"])
        sheet.write(offset - 1, 2, unit, formats["body"])
        if metric == "revenue_growth":
            source_metric = "revenue_growth"
            source_formula = _selected_formula(scenario_rows, source_metric, "I")
            source_cached = forecast[("base", "FY2031", source_metric)]
            terminal_formula = "=TerminalGrowth"
        elif metric == "ebitda_proxy":
            source_formula = f"=D{terminal_row_map['derived_operating_income']}+D{terminal_row_map['depreciation_and_amortization']}"
            source_cached = (
                _decimal(forecast[("base", "FY2031", "derived_operating_income")])
                + _decimal(forecast[("base", "FY2031", "depreciation_and_amortization")])
            )
            terminal_formula = f"=E{terminal_row_map['derived_operating_income']}+E{terminal_row_map['depreciation_and_amortization']}"
        else:
            source_formula = _selected_formula(scenario_rows, metric, "I")
            source_cached = forecast[("base", "FY2031", metric)]
            terminal_formula = ""
        if source_cached is None:
            raise ValueError(f"Missing Base FY2031 source value for terminal bridge: {metric}")
        source_fmt = formats["percent_link"] if unit == "decimal" else formats["amount_link"]
        _write_formula(sheet, offset - 1, 3, source_formula, source_fmt, _number(source_cached))

        if not terminal_formula:
            refs = terminal_row_map
            terminal_formula = {
                "revenue": f"=D{refs['revenue']}*(1+E{refs['revenue_growth']})",
                "gross_margin": f"=D{refs['gross_margin']}",
                "gross_profit": f"=E{refs['revenue']}*E{refs['gross_margin']}",
                "sga_percent_revenue": f"=D{refs['sga_percent_revenue']}",
                "total_selling_and_administrative_expense": f"=E{refs['revenue']}*E{refs['sga_percent_revenue']}",
                "derived_operating_income": f"=E{refs['gross_profit']}-E{refs['total_selling_and_administrative_expense']}",
                "operating_margin": f"=E{refs['derived_operating_income']}/E{refs['revenue']}",
                "normalized_tax_rate": f"=D{refs['normalized_tax_rate']}",
                "normalized_operating_tax_expense": f"=MAX(E{refs['derived_operating_income']},0)*E{refs['normalized_tax_rate']}",
                "nopat": f"=E{refs['derived_operating_income']}-E{refs['normalized_operating_tax_expense']}",
                "da_percent_revenue": f"=D{refs['da_percent_revenue']}",
                "depreciation_and_amortization": f"=E{refs['revenue']}*E{refs['da_percent_revenue']}",
                "capex_percent_revenue": f"=D{refs['capex_percent_revenue']}",
                "capital_expenditures": f"=E{refs['revenue']}*E{refs['capex_percent_revenue']}",
                "operating_nwc_percent_revenue": f"=D{refs['operating_nwc_percent_revenue']}",
                "operating_nwc_proxy": f"=E{refs['revenue']}*E{refs['operating_nwc_percent_revenue']}",
                "change_in_operating_nwc": f"=E{refs['operating_nwc_proxy']}-D{refs['operating_nwc_proxy']}",
                "fcff": f"=E{refs['nopat']}+E{refs['depreciation_and_amortization']}-E{refs['capital_expenditures']}-E{refs['change_in_operating_nwc']}",
            }[metric]
        term_fmt = formats["percent_formula"] if unit == "decimal" else formats["amount_formula"]
        _write_formula(sheet, offset - 1, 4, terminal_formula, term_fmt, _number(cached_terminal))

    sheet.write("B45", "Terminal value", formats["body"])
    _write_formula(sheet, 44, 4, f"=E{terminal_row_map['fcff']}/(CalculatedWACC-TerminalGrowth)", formats["amount_formula"], _number(base_valuation.terminal_value))
    sheet.write("B46", "Present value of terminal value", formats["body"])
    _write_formula(sheet, 45, 4, "=E45/(1+CalculatedWACC)^((DATE(2031,5,31)-ModelDate)/365)", formats["amount_formula"], _number(base_valuation.pv_terminal_value))

    sheet.write("B49", "Enterprise-to-equity bridge", formats["section"])
    for col in range(2, 5):
        sheet.write_blank(48, col, None, formats["section"])
    bridge_rows = {
        "pv_explicit": 50,
        "pv_terminal": 51,
        "enterprise_value": 52,
        "cash": 53,
        "debt": 54,
        "preferred": 55,
        "other": 56,
        "equity_value": 57,
        "diluted_shares": 58,
        "value_per_share": 59,
        "basic_share_value": 60,
        "reference_price": 61,
        "premium_discount": 62,
        "terminal_share": 63,
        "terminal_ev_ebitda": 64,
        "leases": 65,
        "xnpv": 67,
        "explicit_ev": 68,
        "xnpv_diff": 69,
    }
    labels = {
        "pv_explicit": "PV of explicit FCFF",
        "pv_terminal": "PV of terminal value",
        "enterprise_value": "Enterprise value",
        "cash": "Cash and short-term investments",
        "debt": "Less: carrying interest-bearing debt",
        "preferred": "Less: redeemable preferred stock",
        "other": "Other adjustments",
        "equity_value": "Equity value",
        "diluted_shares": "Diluted-proxy shares",
        "value_per_share": "Illustrative value per share",
        "basic_share_value": "Basic-share denominator cross-check",
        "reference_price": "September 4, 2026 reference price",
        "premium_discount": "Premium / (discount) to reference",
        "terminal_share": "PV of terminal value / enterprise value",
        "terminal_ev_ebitda": "Implied terminal EV / EBITDA",
        "leases": "Operating lease liabilities (memorandum)",
        "xnpv": "Native XNPV enterprise value",
        "explicit_ev": "Explicit date-exponent enterprise value",
        "xnpv_diff": "XNPV reconciliation difference",
    }
    bridge_formulas = {
        "pv_explicit": "=I18",
        "pv_terminal": "=E46",
        "enterprise_value": "=E50+E51",
        "cash": "=CashAndShortTermInvestments",
        "debt": "=CarryingDebt",
        "preferred": "=RedeemablePreferredStock",
        "other": "=0",
        "equity_value": "=E52+E53-E54-E55+E56",
        "diluted_shares": "=DilutedProxyShares",
        "value_per_share": "=E57/E58",
        "basic_share_value": "=E57/BasicShares",
        "reference_price": "=ReferenceMarketPrice",
        "premium_discount": "=E59/E61-1",
        "terminal_share": "=E51/E52",
        "terminal_ev_ebitda": f"=E45/E{terminal_row_map['ebitda_proxy']}",
        "leases": "=OperatingLeaseLiabilitiesMemo",
        "xnpv": "=XNPV(CalculatedWACC,D71:I71,D12:I12)",
        "explicit_ev": "=E52",
        "xnpv_diff": "=E67-E68",
    }
    bridge_cached = {
        "pv_explicit": base_valuation.pv_explicit_fcff,
        "pv_terminal": base_valuation.pv_terminal_value,
        "enterprise_value": base_valuation.enterprise_value,
        "cash": model.cash_and_short_term_investments,
        "debt": model.carrying_debt,
        "preferred": model.preferred_stock,
        "other": Decimal("0"),
        "equity_value": base_valuation.equity_value,
        "diluted_shares": model.diluted_proxy_shares,
        "value_per_share": base_valuation.value_per_share,
        "basic_share_value": base_valuation.basic_share_value_cross_check,
        "reference_price": base_valuation.reference_market_price,
        "premium_discount": base_valuation.premium_discount_to_reference,
        "terminal_share": base_valuation.terminal_value_share_of_ev,
        "terminal_ev_ebitda": base_valuation.implied_terminal_ev_ebitda,
        "leases": model.operating_lease_liabilities,
        "xnpv": base_valuation.enterprise_value,
        "explicit_ev": base_valuation.enterprise_value,
        "xnpv_diff": Decimal("0"),
    }
    for key, row in bridge_rows.items():
        sheet.write(row - 1, 1, labels[key], formats["body"])
        if key in {"value_per_share", "basic_share_value", "reference_price"}:
            fmt = formats["total_per_share"] if key == "value_per_share" else formats["per_share_formula"]
        elif key in {"premium_discount", "terminal_share"}:
            fmt = formats["percent_formula"]
        elif key == "terminal_ev_ebitda":
            fmt = formats["multiple"]
        else:
            fmt = formats["total"] if key in {"enterprise_value", "equity_value"} else formats["amount_formula"]
        _write_formula(sheet, row - 1, 4, bridge_formulas[key], fmt, _number(bridge_cached[key]))

    # XNPV helper row: zero on model date and FY2031 FCFF plus terminal value.
    sheet.write("B71", "XNPV cash-flow helper", formats["note"])
    sheet.write_number("D71", 0, formats["amount"])
    for col in range(4, 8):
        letter = xl_col_to_name(col)
        _write_formula(sheet, 70, col, f"={letter}13", formats["amount_formula"], _number(base_valuation.explicit_cash_flows[col - 4].fcff))
    _write_formula(sheet, 70, 8, "=I13+E45", formats["amount_formula"], _number(base_valuation.explicit_cash_flows[-1].fcff + base_valuation.terminal_value))

    # Compact all-scenario formula block used by Cover and reconciliation checks.
    sheet.write("K7", "All-scenario valuation outputs", formats["section"])
    for col in range(11, 14):
        sheet.write_blank(6, col, None, formats["section"])
    sheet.write_row("K8", ["Metric", *SCENARIO_DISPLAY], formats["header"])
    all_rows = {
        "explicit_pv": 9,
        "terminal_revenue": 10,
        "terminal_gross_profit": 11,
        "terminal_sga": 12,
        "terminal_derived": 13,
        "terminal_tax": 14,
        "terminal_nopat": 15,
        "terminal_da": 16,
        "terminal_capex": 17,
        "terminal_onwc": 18,
        "terminal_change_onwc": 19,
        "terminal_fcff": 20,
        "terminal_ebitda": 21,
        "terminal_value": 22,
        "pv_terminal": 23,
        "enterprise_value": 24,
        "equity_value": 25,
        "value_per_share": 26,
        "basic_share_value": 27,
        "premium_discount": 28,
        "terminal_share": 29,
        "terminal_ev_ebitda": 30,
    }
    all_labels = [
        ("explicit_pv", "PV of explicit FCFF"),
        ("terminal_revenue", "FY2032 revenue"),
        ("terminal_gross_profit", "FY2032 gross profit"),
        ("terminal_sga", "FY2032 total SG&A"),
        ("terminal_derived", "FY2032 derived operating income"),
        ("terminal_tax", "FY2032 operating tax expense"),
        ("terminal_nopat", "FY2032 NOPAT"),
        ("terminal_da", "FY2032 D&A"),
        ("terminal_capex", "FY2032 capex"),
        ("terminal_onwc", "FY2032 operating NWC"),
        ("terminal_change_onwc", "FY2032 change in operating NWC"),
        ("terminal_fcff", "FY2032 FCFF"),
        ("terminal_ebitda", "FY2032 EBITDA proxy"),
        ("terminal_value", "Terminal value"),
        ("pv_terminal", "PV of terminal value"),
        ("enterprise_value", "Enterprise value"),
        ("equity_value", "Equity value"),
        ("value_per_share", "Illustrative value per share"),
        ("basic_share_value", "Basic-share cross-check"),
        ("premium_discount", "Premium / (discount)"),
        ("terminal_share", "PV terminal / EV"),
        ("terminal_ev_ebitda", "Terminal EV / EBITDA"),
    ]
    for key, label in all_labels:
        sheet.write(all_rows[key] - 1, 10, label, formats["body"])

    valuation_by_scenario = {item.scenario: item for item in model.scenario_valuations}
    scenario_columns = {"bear": "L", "base": "M", "bull": "N"}
    for scenario, column in scenario_columns.items():
        rows = scenario_rows[scenario]
        val = valuation_by_scenario[scenario]
        term = val.terminal
        formulas = {
            "explicit_pv": f"=SUMPRODUCT('Scenarios'!E{rows['fcff']}:I{rows['fcff']},1/(1+CalculatedWACC)^(($E$12:$I$12-ModelDate)/365))",
            "terminal_revenue": f"='Scenarios'!I{rows['revenue']}*(1+TerminalGrowth)",
            "terminal_gross_profit": f"={column}{all_rows['terminal_revenue']}*'Scenarios'!I{rows['gross_margin']}",
            "terminal_sga": f"={column}{all_rows['terminal_revenue']}*'Scenarios'!I{rows['sga_percent_revenue']}",
            "terminal_derived": f"={column}{all_rows['terminal_gross_profit']}-{column}{all_rows['terminal_sga']}",
            "terminal_tax": f"=MAX({column}{all_rows['terminal_derived']},0)*'Scenarios'!I{rows['normalized_tax_rate']}",
            "terminal_nopat": f"={column}{all_rows['terminal_derived']}-{column}{all_rows['terminal_tax']}",
            "terminal_da": f"={column}{all_rows['terminal_revenue']}*'Scenarios'!I{rows['da_percent_revenue']}",
            "terminal_capex": f"={column}{all_rows['terminal_revenue']}*'Scenarios'!I{rows['capex_percent_revenue']}",
            "terminal_onwc": f"={column}{all_rows['terminal_revenue']}*'Scenarios'!I{rows['operating_nwc_percent_revenue']}",
            "terminal_change_onwc": f"={column}{all_rows['terminal_onwc']}-'Scenarios'!I{rows['operating_nwc_proxy']}",
            "terminal_fcff": f"={column}{all_rows['terminal_nopat']}+{column}{all_rows['terminal_da']}-{column}{all_rows['terminal_capex']}-{column}{all_rows['terminal_change_onwc']}",
            "terminal_ebitda": f"={column}{all_rows['terminal_derived']}+{column}{all_rows['terminal_da']}",
            "terminal_value": f"={column}{all_rows['terminal_fcff']}/(CalculatedWACC-TerminalGrowth)",
            "pv_terminal": f"={column}{all_rows['terminal_value']}/(1+CalculatedWACC)^((DATE(2031,5,31)-ModelDate)/365)",
            "enterprise_value": f"={column}{all_rows['explicit_pv']}+{column}{all_rows['pv_terminal']}",
            "equity_value": f"={column}{all_rows['enterprise_value']}+CashAndShortTermInvestments-CarryingDebt-RedeemablePreferredStock",
            "value_per_share": f"={column}{all_rows['equity_value']}/DilutedProxyShares",
            "basic_share_value": f"={column}{all_rows['equity_value']}/BasicShares",
            "premium_discount": f"={column}{all_rows['value_per_share']}/ReferenceMarketPrice-1",
            "terminal_share": f"={column}{all_rows['pv_terminal']}/{column}{all_rows['enterprise_value']}",
            "terminal_ev_ebitda": f"={column}{all_rows['terminal_value']}/{column}{all_rows['terminal_ebitda']}",
        }
        cached = {
            "explicit_pv": val.pv_explicit_fcff,
            "terminal_revenue": term.revenue,
            "terminal_gross_profit": term.gross_profit,
            "terminal_sga": term.total_selling_and_administrative_expense,
            "terminal_derived": term.derived_operating_income,
            "terminal_tax": term.derived_operating_income - term.nopat,
            "terminal_nopat": term.nopat,
            "terminal_da": term.depreciation_and_amortization,
            "terminal_capex": term.capital_expenditures,
            "terminal_onwc": term.operating_nwc_proxy,
            "terminal_change_onwc": term.change_in_operating_nwc,
            "terminal_fcff": term.fcff,
            "terminal_ebitda": term.ebitda_proxy,
            "terminal_value": val.terminal_value,
            "pv_terminal": val.pv_terminal_value,
            "enterprise_value": val.enterprise_value,
            "equity_value": val.equity_value,
            "value_per_share": val.value_per_share,
            "basic_share_value": val.basic_share_value_cross_check,
            "premium_discount": val.premium_discount_to_reference,
            "terminal_share": val.terminal_value_share_of_ev,
            "terminal_ev_ebitda": val.implied_terminal_ev_ebitda,
        }
        for key, row in all_rows.items():
            if key in {"value_per_share", "basic_share_value"}:
                fmt = formats["per_share_formula"]
            elif key in {"premium_discount", "terminal_share"}:
                fmt = formats["percent_formula"]
            elif key == "terminal_ev_ebitda":
                fmt = formats["multiple"]
            else:
                fmt = formats["amount_formula"]
            _write_formula(sheet, row - 1, ord(column) - ord("A"), formulas[key], fmt, _number(cached[key]))

    workbook.define_name("BearValuePerShare", f"='DCF'!$L${all_rows['value_per_share']}")
    workbook.define_name("BaseValuePerShare", f"='DCF'!$M${all_rows['value_per_share']}")
    workbook.define_name("BullValuePerShare", f"='DCF'!$N${all_rows['value_per_share']}")
    workbook.define_name("SelectedEnterpriseValue", f"='DCF'!$E${bridge_rows['enterprise_value']}")
    workbook.define_name("SelectedEquityValue", f"='DCF'!$E${bridge_rows['equity_value']}")
    workbook.define_name("SelectedValuePerShare", f"='DCF'!$E${bridge_rows['value_per_share']}")
    workbook.define_name("SelectedTerminalValueShare", f"='DCF'!$E${bridge_rows['terminal_share']}")
    # The FY2031 source links and FY2032 terminal formulas intentionally differ
    # from adjacent formulas. Suppress only Excel's inconsistent-formula marker;
    # formula errors remain visible and are also covered by workbook checks.
    sheet.ignore_errors({"formula_differs": "E10:N71"})
    return {
        "bridge_rows": bridge_rows,
        "terminal_rows": terminal_row_map,
        "all_rows": all_rows,
        "scenario_columns": scenario_columns,
    }


def _sensitivity_grid_formula(
    *, base_fcff_row: int, row: int, column: str, helper_fcff_row: int
) -> str:
    """Return the canonical displayed sensitivity formula for one coordinate."""

    rate_cell = f"$C{row}"
    growth_cell = f"{column}$8"
    helper_fcff = f"{column}{helper_fcff_row}"
    return (
        f"=(SUMPRODUCT('Scenarios'!E{base_fcff_row}:I{base_fcff_row},"
        f"1/(1+{rate_cell})^(('DCF'!$E$12:$I$12-ModelDate)/365))"
        f"+({helper_fcff}/({rate_cell}-{growth_cell}))/(1+{rate_cell})^((DATE(2031,5,31)-ModelDate)/365)"
        f"+CashAndShortTermInvestments-CarryingDebt-RedeemablePreferredStock)/DilutedProxyShares"
    )


def _independent_sensitivity_value_formula(
    base_rows: Mapping[str, int], row: int, column: str
) -> str:
    """Recalculate one sensitivity value without using its displayed result/helpers."""

    rate = f"'Sensitivity'!$C${row}"
    growth = f"'Sensitivity'!${column}$8"
    revenue = f"('Scenarios'!$I${base_rows['revenue']}*(1+{growth}))"
    operating_spread = (
        f"('Scenarios'!$I${base_rows['gross_margin']}"
        f"-'Scenarios'!$I${base_rows['sga_percent_revenue']})"
    )
    derived_operating_income = f"({revenue}*{operating_spread})"
    terminal_fcff = (
        f"({derived_operating_income}"
        f"-MAX({derived_operating_income},0)*'Scenarios'!$I${base_rows['normalized_tax_rate']}"
        f"+{revenue}*'Scenarios'!$I${base_rows['da_percent_revenue']}"
        f"-{revenue}*'Scenarios'!$I${base_rows['capex_percent_revenue']}"
        f"-({revenue}*'Scenarios'!$I${base_rows['operating_nwc_percent_revenue']}"
        f"-'Scenarios'!$I${base_rows['operating_nwc_proxy']}))"
    )
    return (
        f"=(SUMPRODUCT('Scenarios'!$E${base_rows['fcff']}:$I${base_rows['fcff']},"
        f"1/(1+{rate})^(('DCF'!$E$12:$I$12-ModelDate)/365))"
        f"+({terminal_fcff}/({rate}-{growth}))/(1+{rate})^((DATE(2031,5,31)-ModelDate)/365)"
        f"+CashAndShortTermInvestments-CarryingDebt-RedeemablePreferredStock)/DilutedProxyShares"
    )


def _write_sensitivity(
    workbook: xlsxwriter.Workbook,
    sheet: xlsxwriter.worksheet.Worksheet,
    formats: Mapping[str, object],
    model: ValuationModel,
    forecast_index: Mapping[tuple[str, str, str], Decimal | None],
    scenario_rows: Mapping[str, Mapping[str, int]],
    input_cells: Mapping[str, str],
) -> dict[str, object]:
    _write_title(sheet, "Base-case sensitivity", "Illustrative value per diluted-proxy share; full-precision WACC and perpetual-growth inputs", formats, 8)
    sheet.hide_gridlines(2)
    sheet.set_zoom(90)
    sheet.freeze_panes(8, 3)
    sheet.set_column("A:A", 2)
    sheet.set_column("B:B", 30)
    sheet.set_column("C:C", 16)
    sheet.set_column("D:H", 15)
    sheet.write("B7", "WACC / perpetual growth", formats["section"])
    for col in range(2, 8):
        sheet.write_blank(6, col, None, formats["section"])
    sheet.write("C8", "WACC", formats["header"])
    growth_inputs = [f"sensitivity_terminal_growth_{bp}bp" for bp in (150, 200, 250, 300, 350)]
    growths = sorted({cell.terminal_growth for cell in model.sensitivity_cells})
    if len(growths) != 5:
        raise ValueError("Sensitivity model must contain exactly five growth inputs.")
    for col, (name, growth) in enumerate(zip(growth_inputs, growths, strict=True), start=3):
        _write_formula(sheet, 7, col, f"={input_cells[name]}", formats["percent_link"], _number(growth))
    delta_names = (
        "sensitivity_wacc_delta_minus_100bp",
        "sensitivity_wacc_delta_minus_50bp",
        "sensitivity_wacc_delta_center",
        "sensitivity_wacc_delta_plus_50bp",
        "sensitivity_wacc_delta_plus_100bp",
    )
    deltas = sorted(
        {cell.wacc - model.wacc_primary.wacc for cell in model.sensitivity_cells}
    )
    if len(deltas) != 5:
        raise ValueError("Sensitivity model must contain exactly five WACC deltas.")
    sensitivity_map = {(cell.wacc, cell.terminal_growth): cell.value_per_share for cell in model.sensitivity_cells}
    for row, (name, delta) in enumerate(zip(delta_names, deltas, strict=True), start=8):
        rate = model.wacc_primary.wacc + delta
        _write_formula(sheet, row, 2, f"=CalculatedWACC+{input_cells[name]}", formats["percent_formula"], _number(rate))

    sheet.write("B17", "Visible terminal-growth helpers", formats["section"])
    for col in range(2, 8):
        sheet.write_blank(16, col, None, formats["section"])
    helper_rows = {
        "growth": 18,
        "revenue": 19,
        "gross_profit": 20,
        "sga": 21,
        "derived": 22,
        "tax": 23,
        "nopat": 24,
        "da": 25,
        "capex": 26,
        "onwc": 27,
        "change_onwc": 28,
        "fcff": 29,
        "ebitda": 30,
    }
    helper_labels = {
        "growth": "Perpetual growth",
        "revenue": "FY2032 revenue",
        "gross_profit": "FY2032 gross profit",
        "sga": "FY2032 total SG&A",
        "derived": "FY2032 derived operating income",
        "tax": "FY2032 tax expense",
        "nopat": "FY2032 NOPAT",
        "da": "FY2032 D&A",
        "capex": "FY2032 capex",
        "onwc": "FY2032 operating NWC",
        "change_onwc": "FY2032 change in operating NWC",
        "fcff": "FY2032 FCFF",
        "ebitda": "FY2032 EBITDA proxy",
    }
    for key, row in helper_rows.items():
        sheet.write(row - 1, 1, helper_labels[key], formats["body"])
    base_forecast = {
        (scenario, year, metric): value
        for (scenario, year, metric), value in forecast_index.items()
        if value is not None
    }
    terminals = [calculate_terminal_bridge("base", base_forecast, growth) for growth in growths]
    base_rows = scenario_rows["base"]
    for col, (growth, terminal) in enumerate(zip(growths, terminals, strict=True), start=3):
        letter = xl_col_to_name(col)
        _write_formula(
            sheet,
            helper_rows["growth"] - 1,
            col,
            f"={letter}$8",
            formats["percent_formula"],
            _number(growth),
        )
        formulas = {
            "revenue": f"='Scenarios'!I{base_rows['revenue']}*(1+{letter}{helper_rows['growth']})",
            "gross_profit": f"={letter}{helper_rows['revenue']}*'Scenarios'!I{base_rows['gross_margin']}",
            "sga": f"={letter}{helper_rows['revenue']}*'Scenarios'!I{base_rows['sga_percent_revenue']}",
            "derived": f"={letter}{helper_rows['gross_profit']}-{letter}{helper_rows['sga']}",
            "tax": f"=MAX({letter}{helper_rows['derived']},0)*'Scenarios'!I{base_rows['normalized_tax_rate']}",
            "nopat": f"={letter}{helper_rows['derived']}-{letter}{helper_rows['tax']}",
            "da": f"={letter}{helper_rows['revenue']}*'Scenarios'!I{base_rows['da_percent_revenue']}",
            "capex": f"={letter}{helper_rows['revenue']}*'Scenarios'!I{base_rows['capex_percent_revenue']}",
            "onwc": f"={letter}{helper_rows['revenue']}*'Scenarios'!I{base_rows['operating_nwc_percent_revenue']}",
            "change_onwc": f"={letter}{helper_rows['onwc']}-'Scenarios'!I{base_rows['operating_nwc_proxy']}",
            "fcff": f"={letter}{helper_rows['nopat']}+{letter}{helper_rows['da']}-{letter}{helper_rows['capex']}-{letter}{helper_rows['change_onwc']}",
            "ebitda": f"={letter}{helper_rows['derived']}+{letter}{helper_rows['da']}",
        }
        cached = {
            "revenue": terminal.revenue,
            "gross_profit": terminal.gross_profit,
            "sga": terminal.total_selling_and_administrative_expense,
            "derived": terminal.derived_operating_income,
            "tax": terminal.derived_operating_income - terminal.nopat,
            "nopat": terminal.nopat,
            "da": terminal.depreciation_and_amortization,
            "capex": terminal.capital_expenditures,
            "onwc": terminal.operating_nwc_proxy,
            "change_onwc": terminal.change_in_operating_nwc,
            "fcff": terminal.fcff,
            "ebitda": terminal.ebitda_proxy,
        }
        for key, formula in formulas.items():
            _write_formula(sheet, helper_rows[key] - 1, col, formula, formats["amount_formula"], _number(cached[key]))

    for row_index, delta in enumerate(deltas, start=8):
        rate = model.wacc_primary.wacc + delta
        for col_index, growth in enumerate(growths, start=3):
            formula = _sensitivity_grid_formula(
                base_fcff_row=base_rows["fcff"],
                row=row_index + 1,
                column=xl_col_to_name(col_index),
                helper_fcff_row=helper_rows["fcff"],
            )
            _write_formula(sheet, row_index, col_index, formula, formats["per_share_formula"], _number(sensitivity_map[(rate, growth)]))
    workbook.define_name("SensitivityCenter", "='Sensitivity'!$F$11")
    sheet.conditional_format("D9:H13", {"type": "3_color_scale", "min_color": "#F4CCCC", "mid_color": "#FFF2CC", "max_color": "#D9EAD3"})
    # The helper block intentionally changes formulas by row (for example, tax uses
    # MAX while adjacent rows do not). Suppress only Excel's inconsistent-formula
    # marker; formula errors remain visible and covered by workbook checks.
    sheet.ignore_errors({"formula_differs": "D19:H30"})
    return {"grid": "D9:H13", "helper_rows": helper_rows}


def _write_cover(
    workbook: xlsxwriter.Workbook,
    sheet: xlsxwriter.worksheet.Worksheet,
    formats: Mapping[str, object],
    model: ValuationModel,
) -> None:
    _write_title(sheet, "Nike valuation model", "Independent portfolio analysis based on reconciled SEC-derived history and approved analyst scenarios", formats, 13)
    sheet.hide_gridlines(2)
    sheet.set_zoom(95)
    sheet.set_column("A:A", 2)
    sheet.set_column("B:B", 34)
    sheet.set_column("C:C", 4)
    sheet.set_column("D:D", 20)
    sheet.set_column("E:F", 3)
    sheet.set_column("G:N", 13)
    sheet.write("B6", "Selected scenario", formats["subheader"])
    sheet.write("D6", "Base", formats["input"])
    sheet.data_validation("D6", {"validate": "list", "source": "=ScenarioList", "input_title": "Scenario", "input_message": "Select Bear, Base, or Bull."})
    workbook.define_name("SelectedScenario", "='Cover'!$D$6")

    metadata = (
        (
            "DCF valuation date",
            "=ModelDate",
            datetime.combine(date.fromisoformat(model.model_date), datetime.min.time()),
            formats["date"],
        ),
        ("Information cutoff", "=InformationCutoff", datetime(2026, 9, 7), formats["date"]),
        ("Reference market date", "=ReferenceMarketDate", datetime(2026, 9, 4), formats["date"]),
    )
    for row, (label, formula, cached, fmt) in enumerate(metadata, start=8):
        sheet.write(row - 1, 1, label, formats["body"])
        _write_formula(sheet, row - 1, 3, formula, fmt, _excel_date_serial(cached))

    sheet.write("B13", "Illustrative value comparison", formats["section"])
    for col in range(2, 4):
        sheet.write_blank(12, col, None, formats["section"])
    sheet.write("B14", "Case", formats["header"])
    sheet.write("D14", "Value per share", formats["header"])
    valuation = {item.scenario: item for item in model.scenario_valuations}
    comparison = (
        ("Bear", "=BearValuePerShare", valuation["bear"].value_per_share),
        ("Base", "=BaseValuePerShare", valuation["base"].value_per_share),
        ("Bull", "=BullValuePerShare", valuation["bull"].value_per_share),
        ("Reference price", "=ReferenceMarketPrice", valuation["base"].reference_market_price),
    )
    for row, (label, formula, cached) in enumerate(comparison, start=15):
        sheet.write(row - 1, 1, label, formats["body"])
        _write_formula(sheet, row - 1, 3, formula, formats["per_share_cover"], _number(cached))

    chart = workbook.add_chart({"type": "column"})
    chart.add_series({
        "name": "Illustrative value per share",
        "categories": "=Cover!$B$15:$B$18",
        "values": "=Cover!$D$15:$D$18",
        "fill": {"color": "#4472C4"},
        "border": {"none": True},
        "points": [
            {"fill": {"color": "#C55A11"}},
            {"fill": {"color": "#4472C4"}},
            {"fill": {"color": "#70AD47"}},
            {"fill": {"color": "#7F7F7F"}},
        ],
        "data_labels": {"value": True, "num_format": "$0.00"},
    })
    chart.set_title({"name": "Illustrative value per share"})
    chart.set_y_axis({"name": "USD per share", "min": 0, "num_format": "$0"})
    chart.set_legend({"none": True})
    chart.set_style(10)
    chart.set_size({"width": 700, "height": 300})
    sheet.insert_chart("G6", chart)

    sheet.write("B21", "Selected-scenario outputs", formats["section"])
    for col in range(2, 4):
        sheet.write_blank(20, col, None, formats["section"])
    selected = (
        ("Enterprise value", "=SelectedEnterpriseValue", valuation["base"].enterprise_value, formats["amount_formula"]),
        ("Equity value", "=SelectedEquityValue", valuation["base"].equity_value, formats["amount_formula"]),
        ("Illustrative value per share", "=SelectedValuePerShare", valuation["base"].value_per_share, formats["per_share_cover"]),
        ("Calculated WACC", "=CalculatedWACC", model.wacc_primary.wacc, formats["percent_audit"]),
        ("Terminal growth", "=TerminalGrowth", Decimal("0.025"), formats["percent"]),
        ("PV terminal value / enterprise value", "=SelectedTerminalValueShare", valuation["base"].terminal_value_share_of_ev, formats["percent"]),
        ("Canonical publication status", "=ModelStatus", "PASS WITH WARNINGS", formats["cover_output"]),
    )
    for row, (label, formula, cached, fmt) in enumerate(selected, start=22):
        sheet.write(row - 1, 1, label, formats["body"])
        _write_formula(sheet, row - 1, 3, formula, fmt, _number(cached) if isinstance(cached, Decimal) else cached)

    sheet.write("B31", "Navigation", formats["section"])
    for col in range(2, 4):
        sheet.write_blank(30, col, None, formats["section"])
    for row, target in enumerate(("Sources", "Historical", "Scenarios", "WACC", "DCF", "Sensitivity", "Checks"), start=32):
        sheet.write_url(row - 1, 1, f"internal:'{target}'!B2", formats["nav"], string=target)
    sheet.merge_range(
        "B41:D43",
        "Annual-model approximation: FY2027 FCFF covers the full fiscal year; Sep. 4 is the discount anchor, "
        "while equity-bridge balances remain at May 31. Checks separate mechanical integrity from reconciliation "
        "to the approved snapshot. This is not an investment recommendation or price target.",
        formats["note"],
    )
    sheet.set_row(40, 24)
    sheet.set_row(41, 24)
    sheet.set_row(42, 24)
    sheet.write("B44", "Formatting legend", formats["subheader"])
    sheet.write("B45", "Editable assumption or control", formats["legend_input"])
    sheet.write("B46", "Same-sheet formula", formats["legend_formula"])
    sheet.write("B47", "Internal cross-sheet reference", formats["legend_link"])


def _build_check_specs(
    model: ValuationModel,
    valuation_assumptions: Sequence[Mapping[str, str]],
    valuation_sources: Sequence[Mapping[str, str]],
    scenario_assumptions: Sequence[Mapping[str, str]],
    forecast_rows: Mapping[tuple[str, str, str], Decimal | None],
    scenario_rows: Mapping[str, Mapping[str, int]],
    wacc_rows: Mapping[str, int],
    dcf_info: Mapping[str, object],
    sensitivity_info: Mapping[str, object],
) -> list[dict[str, object]]:
    specs: list[dict[str, object]] = []

    def add(
        check_id: str,
        category: str,
        scope: str,
        actual: str | float | int,
        expected: object,
        tolerance: float,
        notes: str,
        cached_actual: object | None = None,
        cached_expected: object | None = None,
        text: bool = False,
        control_class: str = "Mechanical integrity",
        difference_formula: str | None = None,
        status_formula: str | None = None,
        cached_status: str | None = None,
    ) -> None:
        resolved_expected = expected if cached_expected is None else cached_expected
        specs.append(
            {
                "check_id": check_id,
                "category": category,
                "scope": scope,
                "actual": actual,
                "expected": expected,
                "tolerance": tolerance,
                "notes": notes,
                "cached_actual": (
                    resolved_expected if cached_actual is None else cached_actual
                ),
                "cached_expected": resolved_expected,
                "text": text,
                "control_class": control_class,
                "difference_formula": difference_formula,
                "status_formula": status_formula,
                "cached_status": cached_status,
            }
        )

    add(
        "desktop_excel_recalculation_sentinel",
        "recalculation",
        "workbook",
        "=1+1",
        2,
        0,
        "Seeded incorrectly before COM; must equal 2 only after desktop Excel recalculation.",
        -999,
        control_class="Package/build control",
    )
    package_build_control = {"control_class": "Package/build control"}
    approved_snapshot = {"control_class": "Approved snapshot"}
    add("valuation_input_count", "inputs", "workbook", "=COUNTA('Sources'!A7:A40)", 34, 0, "Sources!A7:A40", **package_build_control)
    add("valuation_inputs_approved", "inputs", "workbook", '=COUNTIF(\'Sources\'!I7:I40,"approved")', 34, 0, "Sources!I7:I40", **package_build_control)
    add("valuation_source_count", "sources", "workbook", "=COUNTA('Sources'!A45:A55)", 11, 0, "Sources!A45:A55", **package_build_control)
    add("scenario_assumption_count", "inputs", "workbook", "=COUNTA('Sources'!A60:A164)", 105, 0, "Sources!A60:A164", **package_build_control)
    add("scenario_assumptions_approved", "inputs", "workbook", '=COUNTIF(\'Sources\'!J60:J164,"approved")', 105, 0, "Sources!J60:J164", **package_build_control)

    source_ids = [row.get("source_id", "") for row in valuation_sources]
    if not source_ids or any(
        VALUATION_SOURCE_ID_PATTERN.fullmatch(source_id) is None
        for source_id in source_ids
    ):
        raise ValueError(
            "Valuation source IDs must follow the canonical VS plus two ASCII digits grammar."
        )
    source_register_range = "'Sources'!A45:A55"
    add(
        "valuation_source_id_syntax",
        "sources",
        "valuation source register",
        (
            f'=SUMPRODUCT(--(LEN({source_register_range})=4),'
            f'--(LEFT({source_register_range},2)="VS"),'
            f'--ISNUMBER(FIND(MID({source_register_range},3,1),"0123456789")),'
            f'--ISNUMBER(FIND(RIGHT({source_register_range},1),"0123456789")))'
        ),
        len(source_ids),
        0,
        "Every source-register ID follows VS plus exactly two ASCII digits.",
        **package_build_control,
    )
    valuation_source_range = "'Sources'!G7:G40"
    normalized_source_range = (
        f'SUBSTITUTE(SUBSTITUTE(TRIM({valuation_source_range})," ;",";"),"; ",";")'
    )
    valuation_source_expression = (
        f'";"&SUBSTITUTE({normalized_source_range},";",";;")&";"'
    )
    for source_row in range(45, 45 + len(source_ids)):
        valuation_source_expression = (
            f'SUBSTITUTE({valuation_source_expression},'
            f'";"&\'Sources\'!$A${source_row}&";","")'
        )
    add(
        "valuation_source_id_resolution",
        "sources",
        "valuation inputs",
        (
            f'=SUMPRODUCT(--(LEN(TRIM({valuation_source_range}))>0),'
            f'--(LEFT({normalized_source_range},1)<>";"),'
            f'--(RIGHT({normalized_source_range},1)<>";"),'
            f'--ISERROR(SEARCH(";;",{normalized_source_range})),'
            f'--({valuation_source_expression}=""))'
        ),
        34,
        0,
        "Every semicolon-delimited VS## token resolves by exact boundary in tblValuationSources.",
        **package_build_control,
    )
    add(
        "scenario_source_id_resolution",
        "sources",
        "scenario assumptions",
        '=SUMPRODUCT(--ISNUMBER(MATCH(\'Sources\'!H60:H164,{"FS01","FS02","FS03","FS04","FS05","FS06"},0)))',
        105,
        0,
        "Every scenario source ID resolves to the committed Phase 4 forecast-source register.",
        **package_build_control,
    )

    def approved_value_literal(row: Mapping[str, str]) -> str:
        value = row["value"]
        if row.get("unit") == "date":
            parsed = date.fromisoformat(value)
            return f"DATE({parsed.year},{parsed.month},{parsed.day})"
        return value

    valuation_value_checks = ",".join(
        f"'Sources'!D{row_number}={approved_value_literal(row)}"
        for row_number, row in enumerate(valuation_assumptions, start=7)
    )
    scenario_value_checks = ",".join(
        f"'Sources'!D{row_number}={approved_value_literal(row)}"
        for row_number, row in enumerate(scenario_assumptions, start=60)
    )
    add(
        "valuation_input_values_match_approved",
        "inputs",
        "valuation assumptions",
        f"=--AND({valuation_value_checks})",
        1,
        0,
        "Editable valuation values reconcile to the approved committed snapshot.",
        **approved_snapshot,
    )
    add(
        "scenario_input_values_match_approved",
        "inputs",
        "scenario assumptions",
        f"=--AND({scenario_value_checks})",
        1,
        0,
        "Editable scenario values reconcile to the approved committed snapshot.",
        **approved_snapshot,
    )

    add("wacc_weights_sum", "wacc", "primary", f"='WACC'!E{wacc_rows['equity_weight']}+'WACC'!E{wacc_rows['debt_weight']}", 1, 1e-12, "WACC primary weights", 1)
    add("unlevered_beta_formula", "wacc", "primary", f"='WACC'!E{wacc_rows['unlevered_beta']}", _number(model.wacc_primary.unlevered_beta), 1e-9, "WACC product-revenue-weighted beta", **approved_snapshot)
    add("relevered_beta_formula", "wacc", "primary", f"='WACC'!E{wacc_rows['levered_beta']}", _number(model.wacc_primary.levered_beta), 1e-9, "WACC relevering equation", **approved_snapshot)
    add("capm_cost_of_equity", "wacc", "primary", f"='WACC'!E{wacc_rows['cost_equity']}", _number(model.wacc_primary.cost_of_equity), 1e-9, "WACC CAPM equation", **approved_snapshot)
    add("pretax_cost_of_debt", "wacc", "primary", f"='WACC'!E{wacc_rows['pretax_debt']}", _number(model.wacc_primary.pretax_cost_of_debt), 1e-9, "Risk-free rate plus default spread", **approved_snapshot)
    add("after_tax_cost_of_debt", "wacc", "primary", f"='WACC'!E{wacc_rows['after_tax_debt']}", _number(model.wacc_primary.after_tax_cost_of_debt), 1e-9, "Pretax debt cost after operating-tax shield", **approved_snapshot)
    add("wacc_formula", "wacc", "primary", f"='WACC'!E{wacc_rows['wacc']}", _number(model.wacc_primary.wacc), 1e-9, "WACC!E calculated result", **approved_snapshot)
    add("wacc_carrying_crosscheck", "wacc", "carrying debt", f"='WACC'!J{wacc_rows['wacc']}", _number(model.wacc_carrying_cross_check.wacc), 1e-9, "WACC!J calculated result", **approved_snapshot)
    add("wacc_exceeds_terminal_growth", "terminal value", "workbook", "=--(CalculatedWACC>TerminalGrowth)", 1, 0, "Headline WACC must exceed headline terminal growth.")

    for scenario in ("bear", "base", "bull"):
        rows = scenario_rows[scenario]
        for year_offset, year in enumerate(FORECAST_YEARS, start=4):
            col = xl_col_to_name(year_offset)
            for metric in ("revenue", "derived_operating_income", "operating_nwc_proxy", "fcff"):
                expected = forecast_rows[(scenario, year, metric)]
                add(f"forecast_{scenario}_{year}_{metric}", "forecast reconciliation", scenario, f"='Scenarios'!{col}{rows[metric]}", _number(expected), 0.1, f"Scenarios!{col}{rows[metric]}", **approved_snapshot)
            identity_difference = (
                f"='Scenarios'!{col}{rows['fcff']}-("
                f"'Scenarios'!{col}{rows['nopat']}+"
                f"'Scenarios'!{col}{rows['depreciation_and_amortization']}-"
                f"'Scenarios'!{col}{rows['capital_expenditures']}-"
                f"'Scenarios'!{col}{rows['change_in_operating_nwc']})"
            )
            add(
                f"fcff_identity_{scenario}_{year}",
                "FCFF",
                scenario,
                identity_difference,
                0,
                1e-7,
                f"Scenarios!{col}{rows['fcff']} equals NOPAT + D&A - capex - change in operating NWC.",
                cached_actual=0,
            )

    base = next(item for item in model.scenario_valuations if item.scenario == "base")
    for index, flow in enumerate(base.explicit_cash_flows, start=0):
        col = xl_col_to_name(index + 4)
        add(f"discount_exponent_{flow.fiscal_year}", "discounting", "selected scenario", f"='DCF'!{col}14", _number(flow.discount_exponent), 1e-9, f"DCF!{col}14", **approved_snapshot)
        add(f"discount_factor_{flow.fiscal_year}", "discounting", "selected scenario", f"='DCF'!{col}15", _number(flow.discount_factor), 1e-9, f"DCF!{col}15", **approved_snapshot)
    bridge_rows = dcf_info["bridge_rows"]
    terminal_rows = dcf_info["terminal_rows"]
    all_rows = dcf_info["all_rows"]
    scenario_columns = dcf_info["scenario_columns"]

    def selected_expected(metric: str) -> str:
        row = all_rows[metric]
        return (
            "=CHOOSE(MATCH(SelectedScenario,ScenarioList,0),"
            f"'DCF'!{scenario_columns['bear']}{row},"
            f"'DCF'!{scenario_columns['base']}{row},"
            f"'DCF'!{scenario_columns['bull']}{row})"
        )

    add("exact_cash_flow_dates", "discounting", "selected scenario", "=--AND('DCF'!D12=ModelDate,'DCF'!E12=DATE(2027,5,31),'DCF'!F12=DATE(2028,5,31),'DCF'!G12=DATE(2029,5,31),'DCF'!H12=DATE(2030,5,31),'DCF'!I12=DATE(2031,5,31))", 1, 0, "DCF!D12:I12")
    add("xnpv_reconciliation", "discounting", "selected scenario", f"='DCF'!E{bridge_rows['xnpv_diff']}", 0, 0.1, "DCF XNPV and explicit EV")
    add("terminal_bridge_revenue", "terminal bridge", "selected scenario", f"='DCF'!E{terminal_rows['revenue']}", selected_expected("terminal_revenue"), 0.1, "DCF FY2032 bridge", cached_expected=_number(base.terminal.revenue))
    add("terminal_bridge_fcff", "terminal bridge", "selected scenario", f"='DCF'!E{terminal_rows['fcff']}", selected_expected("terminal_fcff"), 0.1, "DCF FY2032 bridge", cached_expected=_number(base.terminal.fcff))
    add("terminal_value_formula", "terminal value", "selected scenario", "='DCF'!E45", selected_expected("terminal_value"), 0.1, "DCF Gordon-growth formula", cached_expected=_number(base.terminal_value))
    add("enterprise_value_equation", "valuation", "selected scenario", f"='DCF'!E{bridge_rows['enterprise_value']}", selected_expected("enterprise_value"), 0.1, "DCF enterprise-to-equity bridge", cached_expected=_number(base.enterprise_value))
    add("equity_value_equation", "valuation", "selected scenario", f"='DCF'!E{bridge_rows['equity_value']}", selected_expected("equity_value"), 0.1, "DCF enterprise-to-equity bridge", cached_expected=_number(base.equity_value))
    add("positive_diluted_shares", "shares", "workbook", "=--(DilutedProxyShares>0)", 1, 0, "Diluted proxy shares must remain positive.")
    add("basic_share_crosscheck", "shares", "selected scenario", f"='DCF'!E{bridge_rows['basic_share_value']}", selected_expected("basic_share_value"), 0.01, "DCF basic-share denominator cross-check", cached_expected=_number(base.basic_share_value_cross_check))

    values = {item.scenario: item for item in model.scenario_valuations}
    for scenario in ("bear", "base", "bull"):
        col = scenario_columns[scenario]
        value = values[scenario]
        add(f"python_value_per_share_{scenario}", "Python reconciliation", scenario, f"='DCF'!{col}{all_rows['value_per_share']}", _number(value.value_per_share), 0.01, "DCF all-scenario formula block", **approved_snapshot)
        add(f"python_enterprise_value_{scenario}", "Python reconciliation", scenario, f"='DCF'!{col}{all_rows['enterprise_value']}", _number(value.enterprise_value), 0.1, "DCF all-scenario formula block", **approved_snapshot)
        add(f"python_terminal_fcff_{scenario}", "terminal bridge", scenario, f"='DCF'!{col}{all_rows['terminal_fcff']}", _number(value.terminal.fcff), 0.1, "DCF all-scenario formula block", **approved_snapshot)
    add("scenario_value_ordering", "scenario coherence", "workbook", f"=--AND('DCF'!{scenario_columns['bear']}{all_rows['value_per_share']}<'DCF'!{scenario_columns['base']}{all_rows['value_per_share']},'DCF'!{scenario_columns['base']}{all_rows['value_per_share']}<'DCF'!{scenario_columns['bull']}{all_rows['value_per_share']})", 1, 0, "Bear < Base < Bull illustrative values")

    grid_rows = list(range(9, 14))
    grid_cols = list("DEFGH")
    base_scenario_rows = scenario_rows["base"]
    ordered_cells = sorted(model.sensitivity_cells, key=lambda cell: (cell.wacc, cell.terminal_growth))
    for row, rate_group in zip(grid_rows, [ordered_cells[i:i + 5] for i in range(0, 25, 5)], strict=True):
        for col, cell in zip(grid_cols, rate_group, strict=True):
            add(
                f"sensitivity_calculation_{col}{row}",
                "sensitivity",
                "Base",
                f"='Sensitivity'!{col}{row}",
                _independent_sensitivity_value_formula(
                    base_scenario_rows, row, col
                ),
                1e-9,
                f"Independent current-coordinate calculation for Sensitivity!{col}{row}.",
                cached_actual=_number(cell.value_per_share),
                cached_expected=_number(cell.value_per_share),
            )
            add(f"sensitivity_{row}_{col}", "sensitivity", "Base", f"='Sensitivity'!{col}{row}", _number(cell.value_per_share), 0.01, f"Sensitivity!{col}{row}", **approved_snapshot)
    center_coordinates_match = (
        "AND(ABS('Sensitivity'!$C$11-CalculatedWACC)<=1E-9,"
        "ABS('Sensitivity'!$F$8-TerminalGrowth)<=1E-9)"
    )
    add(
        "sensitivity_center",
        "sensitivity",
        "Base",
        "=SensitivityCenter",
        "=BaseValuePerShare",
        0.01,
        "Reconcile only when the displayed center coordinates equal the headline WACC and terminal growth.",
        cached_actual=_number(base.value_per_share),
        cached_expected=_number(base.value_per_share),
        control_class="Conditional / input warning",
        difference_formula=f'=IF({center_coordinates_match},D{{row}}-E{{row}},"")',
        status_formula=(
            f'=IF(NOT({center_coordinates_match}),'
            '"NOT APPLICABLE — COORDINATES DIFFER",'
            'IF(AND(ISNUMBER(D{row}),ISNUMBER(E{row}),ISNUMBER(G{row})),'
            'IF(ABS(F{row})<=G{row},"PASS","FAIL"),"UNVERIFIED"))'
        ),
        cached_status="PASS",
    )
    add(
        "sensitivity_growth_coordinate_order",
        "sensitivity coordinates",
        "Base",
        "=--AND('Sensitivity'!D8<'Sensitivity'!E8,'Sensitivity'!E8<'Sensitivity'!F8,'Sensitivity'!F8<'Sensitivity'!G8,'Sensitivity'!G8<'Sensitivity'!H8)",
        1,
        0,
        "Duplicate or unordered growth coordinates are an input warning; cell arithmetic is evaluated separately.",
        control_class="Conditional / input warning",
        status_formula='=IF(AND(ISNUMBER(D{row}),ISNUMBER(E{row})),IF(D{row}=E{row},"PASS","INPUT WARNING"),"UNVERIFIED")',
        cached_status="PASS",
    )
    add(
        "sensitivity_wacc_coordinate_order",
        "sensitivity coordinates",
        "Base",
        "=--AND('Sensitivity'!C9<'Sensitivity'!C10,'Sensitivity'!C10<'Sensitivity'!C11,'Sensitivity'!C11<'Sensitivity'!C12,'Sensitivity'!C12<'Sensitivity'!C13)",
        1,
        0,
        "Duplicate or unordered WACC coordinates are an input warning; cell arithmetic is evaluated separately.",
        control_class="Conditional / input warning",
        status_formula='=IF(AND(ISNUMBER(D{row}),ISNUMBER(E{row})),IF(D{row}=E{row},"PASS","INPUT WARNING"),"UNVERIFIED")',
        cached_status="PASS",
    )
    for row in grid_rows:
        growths_strictly_increase = "AND('Sensitivity'!$D$8<'Sensitivity'!$E$8,'Sensitivity'!$E$8<'Sensitivity'!$F$8,'Sensitivity'!$F$8<'Sensitivity'!$G$8,'Sensitivity'!$G$8<'Sensitivity'!$H$8)"
        add(
            f"sensitivity_growth_monotonicity_{row}",
            "sensitivity",
            "Base",
            f"=--AND('Sensitivity'!D{row}<'Sensitivity'!E{row},'Sensitivity'!E{row}<'Sensitivity'!F{row},'Sensitivity'!F{row}<'Sensitivity'!G{row},'Sensitivity'!G{row}<'Sensitivity'!H{row})",
            1,
            0,
            f"Sensitivity!D{row}:H{row}",
            control_class="Conditional / input warning",
            status_formula=(
                f'=IF(NOT({growths_strictly_increase}),'
                '"NOT APPLICABLE — DUPLICATE OR UNORDERED COORDINATES",'
                'IF(AND(ISNUMBER(D{row}),ISNUMBER(E{row})),'
                'IF(D{row}=E{row},"PASS","FAIL"),"UNVERIFIED"))'
            ),
            cached_status="PASS",
        )
    for col in grid_cols:
        waccs_strictly_increase = "AND('Sensitivity'!$C$9<'Sensitivity'!$C$10,'Sensitivity'!$C$10<'Sensitivity'!$C$11,'Sensitivity'!$C$11<'Sensitivity'!$C$12,'Sensitivity'!$C$12<'Sensitivity'!$C$13)"
        add(
            f"sensitivity_wacc_monotonicity_{col}",
            "sensitivity",
            "Base",
            f"=--AND('Sensitivity'!{col}9>'Sensitivity'!{col}10,'Sensitivity'!{col}10>'Sensitivity'!{col}11,'Sensitivity'!{col}11>'Sensitivity'!{col}12,'Sensitivity'!{col}12>'Sensitivity'!{col}13)",
            1,
            0,
            f"Sensitivity!{col}9:{col}13",
            control_class="Conditional / input warning",
            status_formula=(
                f'=IF(NOT({waccs_strictly_increase}),'
                '"NOT APPLICABLE — DUPLICATE OR UNORDERED COORDINATES",'
                'IF(AND(ISNUMBER(D{row}),ISNUMBER(E{row})),'
                'IF(D{row}=E{row},"PASS","FAIL"),"UNVERIFIED"))'
            ),
            cached_status="PASS",
        )

    error_formula = "=SUMPRODUCT(--ISERROR('Scenarios'!D9:I75))+SUMPRODUCT(--ISERROR('WACC'!E9:E36))+SUMPRODUCT(--ISERROR('WACC'!J9:J36))+SUMPRODUCT(--ISERROR('DCF'!D12:N71))+SUMPRODUCT(--ISERROR('Sensitivity'!C8:H30))"
    add("formula_error_scan", "structure", "workbook", error_formula, 0, 0, "Formula-bearing ranges")
    add("external_workbook_links", "structure", "package", 0, 0, 0, "Verified by package inspection", **package_build_control)
    add("data_connections", "structure", "package", 0, 0, 0, "Verified by package inspection", **package_build_control)
    add("vba_macros", "structure", "package", 0, 0, 0, "Verified by package inspection", **package_build_control)
    return specs


def _write_checks(
    workbook: xlsxwriter.Workbook,
    sheet: xlsxwriter.worksheet.Worksheet,
    formats: Mapping[str, object],
    specs: Sequence[Mapping[str, object]],
    model: ValuationModel,
    dcf_info: Mapping[str, object],
) -> int:
    _write_title(sheet, "Model checks", "Mechanical integrity, approved-snapshot reconciliation, warnings, and package/build controls remain distinct", formats, 9)
    sheet.hide_gridlines(2)
    sheet.set_zoom(80)
    sheet.freeze_panes(8, 3)
    sheet.set_column("A:A", 48)
    sheet.set_column("B:B", 28)
    sheet.set_column("C:C", 32)
    sheet.set_column("D:F", 18)
    sheet.set_column("G:G", 26)
    sheet.set_column("H:H", 58)
    sheet.set_column("I:I", 56)
    sheet.write("A5", "Mechanical integrity", formats["subheader"])
    sheet.write("C5", "Approved snapshot", formats["subheader"])
    sheet.write("E5", "Warnings", formats["subheader"])
    sheet.write("G5", "Package/build controls", formats["subheader"])
    sheet.write("H5", "Canonical publication gate", formats["subheader"])
    sheet.write("A8", "Check ID", formats["header"])
    sheet.write_row("B8", ["Control class", "Category / scope", "Actual", "Expected", "Difference", "Tolerance", "Status", "Notes / fix location"], formats["header"])
    start_row = 9
    for offset, spec in enumerate(specs, start=start_row):
        sheet.write(offset - 1, 0, spec["check_id"], formats["body"])
        sheet.write(offset - 1, 1, spec["control_class"], formats["body"])
        sheet.write(
            offset - 1,
            2,
            f"{spec['category']} — {spec['scope']}",
            formats["body"],
        )
        actual = spec["actual"]
        if isinstance(actual, str) and actual.startswith("="):
            _write_formula(sheet, offset - 1, 3, actual, formats["formula"], spec["cached_actual"])
        else:
            sheet.write(offset - 1, 3, actual, formats["source"])
        expected = spec["expected"]
        if isinstance(expected, str) and expected.startswith("="):
            _write_formula(
                sheet,
                offset - 1,
                4,
                expected,
                formats["formula"],
                spec["cached_expected"],
            )
        else:
            sheet.write(offset - 1, 4, expected, formats["source"])
        difference_formula = spec["difference_formula"]
        if difference_formula is None:
            difference_formula = f"=D{offset}-E{offset}"
            cached_difference: object = _number(spec["cached_actual"]) - _number(
                spec["cached_expected"]
            )
        else:
            difference_formula = str(difference_formula).format(row=offset)
            cached_difference = _number(spec["cached_actual"]) - _number(
                spec["cached_expected"]
            )
        _write_formula(
            sheet,
            offset - 1,
            5,
            difference_formula,
            formats["formula"],
            cached_difference,
        )
        sheet.write_number(offset - 1, 6, float(spec["tolerance"]), formats["source"])
        status_formula = spec["status_formula"]
        if status_formula is None:
            success = (
                "MATCHES APPROVED"
                if spec["control_class"] == "Approved snapshot"
                else "PASS"
            )
            failure = (
                "DIFFERS FROM APPROVED"
                if spec["control_class"] == "Approved snapshot"
                else "FAIL"
            )
            status_formula = (
                f'=IF(AND(ISNUMBER(D{offset}),ISNUMBER(E{offset}),ISNUMBER(F{offset}),ISNUMBER(G{offset})),'
                f'IF(ABS(F{offset})<=G{offset},"{success}","{failure}"),"UNVERIFIED")'
            )
            cached_status = success
        else:
            status_formula = str(status_formula).format(row=offset)
            cached_status = spec["cached_status"] or "UNVERIFIED"
        _write_formula(
            sheet,
            offset - 1,
            7,
            status_formula,
            formats["formula"],
            cached_status,
        )
        sheet.write(offset - 1, 8, spec["notes"], formats["source_wrap"])

    terminal_start = start_row + len(specs)
    all_rows = dcf_info["all_rows"]
    scenario_columns = dcf_info["scenario_columns"]
    for index, scenario in enumerate(("bear", "base", "bull"), start=0):
        row = terminal_start + index
        col = scenario_columns[scenario]
        share = next(item for item in model.scenario_valuations if item.scenario == scenario).terminal_value_share_of_ev
        sheet.write(row - 1, 0, f"terminal_value_dependence_{scenario}", formats["body"])
        sheet.write(row - 1, 1, "Valuation warning", formats["body"])
        sheet.write(row - 1, 2, f"terminal value — {scenario}", formats["body"])
        _write_formula(sheet, row - 1, 3, f"='DCF'!{col}{all_rows['terminal_share']}", formats["percent_formula"], _number(share))
        sheet.write(row - 1, 4, "warning >75%; strong warning >80%", formats["source_wrap"])
        sheet.write_blank(row - 1, 5, None, formats["body"])
        sheet.write_blank(row - 1, 6, None, formats["body"])
        status = "STRONG WARNING" if share > Decimal("0.80") else "WARNING" if share > Decimal("0.75") else "PASS"
        _write_formula(sheet, row - 1, 7, f'=IF(D{row}>0.8,"STRONG WARNING",IF(D{row}>0.75,"WARNING","PASS"))', formats["warning"], status)
        sheet.write(row - 1, 8, "Dependence observation; does not block valuation output.", formats["source_wrap"])

    last_row = terminal_start + 2
    class_range = f"B{start_row}:B{last_row}"
    status_range = f"H{start_row}:H{last_row}"
    mechanical_status = (
        f'=IF(COUNTIF({class_range},"Mechanical integrity")=0,"UNVERIFIED",'
        f'IF(COUNTIFS({class_range},"Mechanical integrity",{status_range},"FAIL")>0,"FAIL",'
        f'IF(COUNTIFS({class_range},"Mechanical integrity",{status_range},"<>PASS")>0,"UNVERIFIED","PASS")))'
    )
    snapshot_status = (
        f'=IF(COUNTIF({class_range},"Approved snapshot")=0,"UNVERIFIED",'
        f'IF(COUNTIFS({class_range},"Approved snapshot",{status_range},"DIFFERS FROM APPROVED")>0,"DIFFERS FROM APPROVED",'
        f'IF(COUNTIFS({class_range},"Approved snapshot",{status_range},"<>MATCHES APPROVED")>0,"UNVERIFIED","MATCHES APPROVED")))'
    )
    warning_status = (
        f'=IF(COUNTIF({class_range},"*warning")=0,"UNVERIFIED",'
        f'IF(COUNTIFS({class_range},"*warning",{status_range},"*WARNING*")>0,"WARNINGS PRESENT",'
        f'IF(COUNTIFS({class_range},"*warning",{status_range},"<>PASS")>0,"UNVERIFIED","PASS")))'
    )
    package_build_control_status = (
        f'=IF(COUNTIF({class_range},"Package/build control")=0,"UNVERIFIED",'
        f'IF(COUNTIFS({class_range},"Package/build control",{status_range},"FAIL")>0,"FAIL",'
        f'IF(COUNTIFS({class_range},"Package/build control",{status_range},"<>PASS")>0,"UNVERIFIED","PASS")))'
    )
    model_status_formula = (
        '=IF(OR(A6="FAIL",G6="FAIL"),"FAIL",'
        'IF(OR(A6<>"PASS",G6<>"PASS"),"UNVERIFIED",'
        'IF(C6="DIFFERS FROM APPROVED","NOT PUBLISHABLE — DIFFERS FROM APPROVED",'
        'IF(C6<>"MATCHES APPROVED","UNVERIFIED",'
        'IF(E6="WARNINGS PRESENT","PASS WITH WARNINGS",'
        'IF(E6="PASS","PASS","UNVERIFIED"))))))'
    )
    _write_formula(sheet, 5, 0, mechanical_status, formats["pass"], "PASS")
    _write_formula(sheet, 5, 2, snapshot_status, formats["pass"], "MATCHES APPROVED")
    _write_formula(sheet, 5, 4, warning_status, formats["warning"], "WARNINGS PRESENT")
    _write_formula(sheet, 5, 6, package_build_control_status, formats["pass"], "PASS")
    _write_formula(sheet, 5, 7, model_status_formula, formats["warning"], "PASS WITH WARNINGS")
    workbook.define_name("MechanicalIntegrityStatus", "='Checks'!$A$6")
    workbook.define_name("ApprovedSnapshotStatus", "='Checks'!$C$6")
    workbook.define_name("WarningStatus", "='Checks'!$E$6")
    workbook.define_name("PackageBuildControlStatus", "='Checks'!$G$6")
    workbook.define_name("ModelStatus", "='Checks'!$H$6")
    sheet.conditional_format(f"H{start_row}:H{last_row}", {"type": "text", "criteria": "containing", "value": "FAIL", "format": formats["fail"]})
    sheet.conditional_format(f"H{start_row}:H{last_row}", {"type": "text", "criteria": "containing", "value": "DIFFERS", "format": formats["warning"]})
    sheet.conditional_format(f"H{start_row}:H{last_row}", {"type": "text", "criteria": "containing", "value": "NOT APPLICABLE", "format": formats["warning"]})
    sheet.conditional_format(f"H{start_row}:H{last_row}", {"type": "text", "criteria": "containing", "value": "UNVERIFIED", "format": formats["fail"]})
    sheet.conditional_format(f"H{start_row}:H{last_row}", {"type": "text", "criteria": "containing", "value": "WARNING", "format": formats["warning"]})
    sheet.conditional_format(f"H{start_row}:H{last_row}", {"type": "text", "criteria": "containing", "value": "PASS", "format": formats["pass"]})
    return last_row


def build_workbook(repository_root: Path, output_path: Path) -> dict[str, object]:
    """Build one formula-driven workbook candidate from committed local inputs."""

    verify_all_protected_artifacts(repository_root)
    model = build_valuation_model(repository_root)
    assumptions = _read_csv(repository_root / "config/valuation_assumptions.csv")
    sources = _read_csv(repository_root / "config/valuation_sources.csv")
    scenario_assumptions = _read_csv(repository_root / "config/scenario_assumptions.csv")
    forecast_rows_raw = _read_csv(repository_root / "outputs/tables/scenario_forecast_long.csv")
    forecast = _forecast_index(forecast_rows_raw)
    assumptions_by_name = {row["input_name"]: row for row in assumptions}
    if len(assumptions) != 34 or any(row["owner_status"] != "approved" for row in assumptions):
        raise ValueError("Valuation assumptions must contain exactly 34 approved rows.")
    if len(scenario_assumptions) != 105 or any(row["owner_status"] != "approved" for row in scenario_assumptions):
        raise ValueError("Scenario assumptions must contain exactly 105 approved rows.")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = xlsxwriter.Workbook(output_path, {"strings_to_urls": False, "nan_inf_to_errors": True})
    workbook.set_properties({
        "title": "Nike valuation model",
        "subject": "Formula-driven DCF translation of approved Phase 5A methodology",
        "author": MODEL_AUTHOR,
        "company": "",
        "comments": "Independent portfolio analysis; not an investment recommendation.",
        "created": datetime(2026, 5, 31),
    })
    # Manual mode prevents Excel from starting an automatic background rebuild while
    # the fail-closed recalculation script is opening the temporary candidate. The
    # verified package is normalized back to automatic/full-recalculation settings.
    workbook.set_calc_mode("manual")
    formats = _format_catalog(workbook)
    sheets = {name: workbook.add_worksheet(name) for name in SHEET_ORDER}
    input_cells, scenario_cells = _write_sources(workbook, sheets["Sources"], formats, assumptions, sources, scenario_assumptions, model)
    _write_historical(sheets["Historical"], formats, forecast, assumptions_by_name, model)
    scenario_rows = _write_scenarios(sheets["Scenarios"], formats, forecast, scenario_cells)
    wacc_rows = _write_wacc(workbook, sheets["WACC"], formats, model)
    dcf_info = _write_dcf(workbook, sheets["DCF"], formats, model, forecast, scenario_rows)
    sensitivity_info = _write_sensitivity(workbook, sheets["Sensitivity"], formats, model, forecast, scenario_rows, input_cells)
    check_specs = _build_check_specs(
        model,
        assumptions,
        sources,
        scenario_assumptions,
        forecast,
        scenario_rows,
        wacc_rows,
        dcf_info,
        sensitivity_info,
    )
    check_rows = _write_checks(workbook, sheets["Checks"], formats, check_specs, model, dcf_info)
    _write_cover(workbook, sheets["Cover"], formats, model)
    workbook.close()
    return {"model": model, "checks": check_rows - 8, "scenario_rows": scenario_rows}


def normalize_workbook_package(path: Path) -> None:
    """Normalize volatile OOXML package metadata without changing formula caches."""

    replacement = path.with_suffix(".normalized.xlsx")
    with ZipFile(path, "r") as source, ZipFile(replacement, "w", ZIP_DEFLATED, compresslevel=9) as target:
        for original in source.infolist():
            data = source.read(original.filename)
            if original.filename == "docProps/core.xml":
                root = ET.fromstring(data)
                namespaces = {
                    "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
                    "dc": "http://purl.org/dc/elements/1.1/",
                    "dcterms": "http://purl.org/dc/terms/",
                }
                creator = root.find("dc:creator", namespaces)
                if creator is not None:
                    creator.text = MODEL_AUTHOR
                modified_by = root.find("cp:lastModifiedBy", namespaces)
                if modified_by is not None:
                    modified_by.text = MODEL_AUTHOR
                for tag in ("created", "modified"):
                    element = root.find(f"dcterms:{tag}", namespaces)
                    if element is not None:
                        element.text = FIXED_CORE_TIMESTAMP
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            elif original.filename == "xl/workbook.xml":
                workbook_xml = data.decode("utf-8")
                workbook_xml = re.sub(
                    r"<ext\b[^>]*>\s*<(?:[A-Za-z_][\w.-]*:)?absPath\b[^>]*/>\s*</ext>",
                    "",
                    workbook_xml,
                )
                workbook_xml = re.sub(
                    r"<(?:[A-Za-z_][\w.-]*:)?absPath\b[^>]*/>",
                    "",
                    workbook_xml,
                )
                workbook_xml = re.sub(
                    r"<extLst\b[^>]*>\s*</extLst>", "", workbook_xml
                )
                calculation_match = re.search(r"<calcPr\b[^>]*/>", workbook_xml)
                if calculation_match is None:
                    raise ValueError("Workbook package is missing calculation settings.")
                calculation_tag = calculation_match.group(0)
                for attribute, value in (
                    ("calcMode", "auto"),
                    ("fullCalcOnLoad", "1"),
                    ("forceFullCalc", "1"),
                ):
                    pattern = rf'\s{attribute}="[^"]*"'
                    if re.search(pattern, calculation_tag):
                        calculation_tag = re.sub(
                            pattern, f' {attribute}="{value}"', calculation_tag
                        )
                    else:
                        calculation_tag = calculation_tag.replace(
                            "/>", f' {attribute}="{value}"/>'
                        )
                workbook_xml = (
                    workbook_xml[: calculation_match.start()]
                    + calculation_tag
                    + workbook_xml[calculation_match.end() :]
                )
                data = workbook_xml.encode("utf-8")
            info = ZipInfo(original.filename, FIXED_PACKAGE_TIMESTAMP)
            info.compress_type = ZIP_DEFLATED
            info.create_system = 0
            info.external_attr = 0
            info.flag_bits = original.flag_bits
            target.writestr(info, data)
    os.replace(replacement, path)


def _defined_name_set(workbook: openpyxl.Workbook) -> set[str]:
    return set(workbook.defined_names)


def _package_scan(path: Path) -> dict[str, int]:
    with ZipFile(path, "r") as archive:
        names = archive.namelist()
        text_parts: list[str] = []
        for name in names:
            if name.endswith((".xml", ".rels")):
                try:
                    text_parts.append(archive.read(name).decode("utf-8", errors="ignore"))
                except KeyError:
                    continue
    package_text = "\n".join(text_parts)
    privacy_patterns = (
        r"C:\\Users\\",
        r"SEC_USER_AGENT",
        r"[A-Za-z0-9._%+-]+@(gmail|yahoo|hotmail|outlook)\.[A-Za-z]{2,}",
    )
    if any(re.search(pattern, package_text, flags=re.IGNORECASE) for pattern in privacy_patterns):
        raise ValueError("Workbook package contains a prohibited private or machine-specific value.")
    return {
        "external_links": sum(name.startswith("xl/externalLinks/") for name in names),
        "connections": sum(name == "xl/connections.xml" for name in names),
        "vba": sum(name.endswith("vbaProject.bin") for name in names),
    }


def _cell_number(workbook: openpyxl.Workbook, sheet: str, address: str) -> Decimal:
    value = workbook[sheet][address].value
    if value is None or isinstance(value, str):
        raise ValueError(f"Required recalculated cell is missing or nonnumeric: {sheet}!{address}")
    return Decimal(str(value))


def _embedded_value_matches(actual: object, expected: str, unit: str = "") -> bool:
    """Compare one embedded workbook value with its committed CSV representation."""

    if expected == "":
        return actual is None
    if unit == "date" or re.fullmatch(r"\d{4}-\d{2}-\d{2}", expected):
        if not isinstance(actual, (date, datetime)):
            return False
        actual_date = actual.date() if isinstance(actual, datetime) else actual
        return actual_date.isoformat() == expected
    if unit in {"decimal", "USD millions", "USD per share", "millions", "x"}:
        if actual is None or isinstance(actual, str):
            return False
        return abs(Decimal(str(actual)) - Decimal(expected)) <= Decimal("1E-9")
    return actual == expected


def _assert_embedded_inputs_match(
    workbook: openpyxl.Workbook, repository_root: Path
) -> None:
    """Reject a publication candidate whose embedded inputs differ from Git inputs."""

    sheet = workbook["Sources"]
    tables = (
        (
            _read_csv(repository_root / "config/valuation_assumptions.csv"),
            7,
            lambda row, field: (
                row.get("unit", "")
                if field == "value"
                else "date"
                if field == "observation_date" and row.get(field)
                else ""
            ),
        ),
        (
            _read_csv(repository_root / "config/valuation_sources.csv"),
            45,
            lambda row, field: (
                "date"
                if field
                in {
                    "publication_date",
                    "observation_date",
                    "information_cutoff",
                    "retrieval_date",
                }
                and row.get(field)
                else ""
            ),
        ),
        (
            _read_csv(repository_root / "config/scenario_assumptions.csv"),
            60,
            lambda row, field: row.get("unit", "") if field == "value" else "",
        ),
    )
    for rows, start_row, unit_for in tables:
        if not rows:
            raise ValueError("A required committed input register is empty.")
        fields = list(rows[0])
        for row_number, row in enumerate(rows, start=start_row):
            for column_number, field in enumerate(fields, start=1):
                if not _embedded_value_matches(
                    sheet.cell(row_number, column_number).value,
                    row[field],
                    unit_for(row, field),
                ):
                    raise ValueError(
                        "Workbook embedded inputs differ from the approved committed snapshot."
                    )


def _normalized_sensitivity_formula(formula: object) -> str:
    """Normalize only Excel's harmless quote removal around simple sheet names."""

    return str(formula).replace("'Scenarios'!", "Scenarios!").replace(
        "'DCF'!", "DCF!"
    )


def _assert_sensitivity_center_formula_contract(workbook: openpyxl.Workbook) -> None:
    """Reject the reproduced same-value replacement of the canonical center formula."""

    base_fcff_row = SCENARIO_BLOCK_STARTS["base"] + next(
        offset
        for offset, (metric, _label, _unit, _driver) in enumerate(
            SCENARIO_METRICS, start=2
        )
        if metric == "fcff"
    )
    expected = _sensitivity_grid_formula(
        base_fcff_row=base_fcff_row,
        row=11,
        column="F",
        helper_fcff_row=29,
    )
    actual = workbook["Sensitivity"]["F11"].value
    if _normalized_sensitivity_formula(actual) != _normalized_sensitivity_formula(
        expected
    ):
        raise ValueError(
            "Sensitivity formula-integrity mismatch at Sensitivity!F11."
        )


def inspect_workbook(
    path: Path,
    repository_root: Path,
    *,
    require_recalculated: bool = True,
    expected_scenario: str = "Base",
) -> dict[str, object]:
    """Inspect workbook structure, formulas, caches, reconciliation, and privacy."""

    if expected_scenario not in SCENARIO_DISPLAY:
        raise ValueError("Expected scenario must be Bear, Base, or Bull.")
    formula_book = openpyxl.load_workbook(path, data_only=False, read_only=False, keep_links=True)
    data_book = openpyxl.load_workbook(path, data_only=True, read_only=False, keep_links=True)
    try:
        if tuple(formula_book.sheetnames) != SHEET_ORDER:
            raise ValueError("Workbook sheet order does not match the approved eight-sheet design.")
        if any(formula_book[name].sheet_state != "visible" for name in SHEET_ORDER):
            raise ValueError("Every workbook sheet must remain visible.")
        names = _defined_name_set(formula_book)
        missing_names = REQUIRED_DEFINED_NAMES - names
        if missing_names:
            raise ValueError("Workbook is missing required defined names.")
        table_names = {table.name for sheet in formula_book.worksheets for table in sheet.tables.values()}
        if table_names != {"tblValuationInputs", "tblValuationSources", "tblScenarioAssumptions"}:
            raise ValueError("Workbook source-table names do not match the approved design.")
        validations = list(formula_book["Cover"].data_validations.dataValidation)
        if len(validations) != 1 or "D6" not in str(validations[0].sqref):
            raise ValueError("Cover!D6 must contain the sole scenario data-validation control.")
        if formula_book["Cover"]["D6"].value != expected_scenario:
            raise ValueError(f"Cover!D6 must contain {expected_scenario}.")
        _assert_embedded_inputs_match(formula_book, repository_root)
        _assert_sensitivity_center_formula_contract(formula_book)
        expected_navigation = ("Sources", "Historical", "Scenarios", "WACC", "DCF", "Sensitivity", "Checks")
        for row, target in enumerate(expected_navigation, start=32):
            link = formula_book["Cover"][f"B{row}"].hyperlink
            if link is None or link.location != f"'{target}'!B2":
                raise ValueError("Cover navigation links do not match the approved visible sheets.")
        cover_charts = formula_book["Cover"]._charts
        if len(cover_charts) != 1 or len(cover_charts[0].ser) != 1:
            raise ValueError("Cover must contain the single approved value-comparison chart.")
        cover_series = cover_charts[0].ser[0]
        if (
            cover_series.cat is None
            or cover_series.cat.strRef is None
            or cover_series.cat.strRef.f != "Cover!$B$15:$B$18"
            or cover_series.val is None
            or cover_series.val.numRef is None
            or cover_series.val.numRef.f != "Cover!$D$15:$D$18"
        ):
            raise ValueError("Cover chart does not reference the approved scenario and reference-price range.")
        calculation = formula_book.calculation
        if require_recalculated and (
            calculation.calcMode != "auto"
            or not calculation.fullCalcOnLoad
            or not calculation.forceFullCalc
        ):
            raise ValueError(
                "Workbook calculation settings must request automatic and full recalculation."
            )
        formulas: list[tuple[str, str, str]] = []
        formula_errors: list[str] = []
        missing_caches: list[str] = []
        for sheet in formula_book.worksheets:
            data_sheet = data_book[sheet.title]
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        formula = str(cell.value)
                        formulas.append((sheet.title, cell.coordinate, formula))
                        if re.search(r"\b(INDIRECT|OFFSET)\s*\(", formula, re.IGNORECASE):
                            raise ValueError("Workbook contains a prohibited volatile selection formula.")
                        cached = data_sheet[cell.coordinate].value
                        if require_recalculated and cached is None:
                            missing_caches.append(f"{sheet.title}!{cell.coordinate}")
                        if isinstance(cached, str) and cached.upper() in FORMULA_ERRORS:
                            formula_errors.append(f"{sheet.title}!{cell.coordinate}")
        if formula_errors:
            raise ValueError("Workbook contains recalculated formula errors.")
        if missing_caches:
            raise ValueError("Workbook contains stale or missing required formula caches.")
        if not any("XNPV(" in formula.upper() for _, _, formula in formulas):
            raise ValueError("Workbook is missing the required native XNPV formula.")
        package = _package_scan(path)
        if any(package.values()):
            raise ValueError("Workbook contains a prohibited external link, connection, or VBA package part.")

        model = build_valuation_model(repository_root)
        wacc = _cell_number(data_book, "WACC", "E36")
        carrying_wacc = _cell_number(data_book, "WACC", "J36")
        if abs(wacc - model.wacc_primary.wacc) > Decimal("1E-9"):
            raise ValueError("Excel primary WACC does not reconcile to Python.")
        if abs(carrying_wacc - model.wacc_carrying_cross_check.wacc) > Decimal("1E-9"):
            raise ValueError("Excel carrying-debt WACC does not reconcile to Python.")
        scenario_cells = {"bear": "L26", "base": "M26", "bull": "N26"}
        reconciliations: dict[str, str] = {}
        for valuation in model.scenario_valuations:
            excel_value = _cell_number(data_book, "DCF", scenario_cells[valuation.scenario])
            difference = abs(excel_value - valuation.value_per_share)
            if difference > Decimal("0.01"):
                raise ValueError(f"Excel {valuation.scenario} value per share does not reconcile to Python.")
            reconciliations[valuation.scenario] = format(difference, "f")
        xnpv_difference = abs(_cell_number(data_book, "DCF", "E69"))
        if xnpv_difference > Decimal("0.1"):
            raise ValueError("Native XNPV does not reconcile to explicit date-exponent PV.")
        center_wacc_difference = abs(
            _cell_number(data_book, "Sensitivity", "C11")
            - _cell_number(data_book, "WACC", "E36")
        )
        center_growth_difference = abs(
            _cell_number(data_book, "Sensitivity", "F8")
            - _cell_number(data_book, "Sources", "D30")
        )
        if (
            center_wacc_difference > Decimal("1E-9")
            or center_growth_difference > Decimal("1E-9")
        ):
            raise ValueError(
                "Canonical workbook sensitivity-center coordinates do not match headline assumptions."
            )
        center_difference = abs(_cell_number(data_book, "Sensitivity", "F11") - next(item for item in model.scenario_valuations if item.scenario == "base").value_per_share)
        if center_difference > Decimal("0.01"):
            raise ValueError("Sensitivity center does not reconcile to Base value per share.")
        for row in range(9, 14):
            values = [_cell_number(data_book, "Sensitivity", f"{col}{row}") for col in "DEFGH"]
            if not all(left < right for left, right in zip(values, values[1:])):
                raise ValueError("Sensitivity is not increasing with perpetual growth.")
        for col in "DEFGH":
            values = [_cell_number(data_book, "Sensitivity", f"{col}{row}") for row in range(9, 14)]
            if not all(left > right for left, right in zip(values, values[1:])):
                raise ValueError("Sensitivity is not decreasing with WACC.")
        model_status = data_book["Checks"]["H6"].value
        if model_status != "PASS WITH WARNINGS":
            raise ValueError(
                "Workbook canonical publication status is not PASS WITH WARNINGS."
            )
        check_sheet = formula_book["Checks"]
        sentinel_rows = [
            row
            for row in range(9, check_sheet.max_row + 1)
            if check_sheet[f"A{row}"].value == "desktop_excel_recalculation_sentinel"
        ]
        if require_recalculated and (
            sentinel_rows != [9]
            or _cell_number(data_book, "Checks", "D9") != Decimal("2")
        ):
            raise ValueError(
                "Desktop Excel recalculation sentinel did not replace its invalid seed."
            )
        control_rows = [
            (
                data_book["Checks"][f"B{row}"].value,
                data_book["Checks"][f"H{row}"].value,
            )
            for row in range(9, data_book["Checks"].max_row + 1)
        ]
        statuses = [status for _, status in control_rows]
        expected_summaries = {
            "A6": "PASS",
            "C6": "MATCHES APPROVED",
            "E6": "WARNINGS PRESENT",
            "G6": "PASS",
        }
        if any(
            data_book["Checks"][cell].value != expected
            for cell, expected in expected_summaries.items()
        ):
            raise ValueError("Workbook control-class summaries are not canonical.")
        allowed_statuses = {
            "Mechanical integrity": {"PASS"},
            "Approved snapshot": {"MATCHES APPROVED"},
            "Package/build control": {"PASS"},
            "Conditional / input warning": {"PASS"},
            "Valuation warning": {"PASS", "WARNING", "STRONG WARNING"},
        }
        if any(
            control_class not in allowed_statuses
            or status not in allowed_statuses[control_class]
            for control_class, status in control_rows
        ):
            raise ValueError(
                "Workbook contains a noncanonical, failed, unverified, or inapplicable control."
            )
        return {
            "sheets": tuple(formula_book.sheetnames),
            "formula_count": len(formulas),
            "check_status": model_status,
            "pass_count": statuses.count("PASS"),
            "snapshot_match_count": statuses.count("MATCHES APPROVED"),
            "warning_count": sum(
                status in {"WARNING", "STRONG WARNING", "INPUT WARNING"}
                for status in statuses
            ),
            "fail_count": statuses.count("FAIL"),
            "mechanical_status": data_book["Checks"]["A6"].value,
            "snapshot_status": data_book["Checks"]["C6"].value,
            "warning_status": data_book["Checks"]["E6"].value,
            "package_build_control_status": data_book["Checks"]["G6"].value,
            "xnpv_difference": format(xnpv_difference, "f"),
            "scenario_differences": reconciliations,
            "package": package,
        }
    finally:
        formula_book.close()
        data_book.close()


def semantic_digest(path: Path) -> str:
    """Return a stable digest of workbook structure, formulas, results, and formats."""

    formula_book = openpyxl.load_workbook(path, data_only=False, read_only=False, keep_links=True)
    data_book = openpyxl.load_workbook(path, data_only=True, read_only=False, keep_links=True)
    try:
        payload: dict[str, object] = {"sheets": [], "names": [], "package": _package_scan(path)}
        for name in formula_book.defined_names:
            definition = formula_book.defined_names[name]
            payload["names"].append((name, definition.attr_text))
        for sheet in formula_book.worksheets:
            data_sheet = data_book[sheet.title]
            cells = []
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    cells.append((cell.coordinate, cell.value, data_sheet[cell.coordinate].value, cell.number_format, cell.style_id))
            tables = sorted((table.name, table.ref) for table in sheet.tables.values())
            validations = sorted((str(item.sqref), item.type, item.formula1) for item in sheet.data_validations.dataValidation)
            charts = []
            for chart in sheet._charts:
                series = []
                for item in chart.series:
                    categories = getattr(getattr(item, "cat", None), "strRef", None) or getattr(getattr(item, "cat", None), "numRef", None)
                    values = getattr(getattr(item, "val", None), "numRef", None)
                    series.append((getattr(categories, "f", None), getattr(values, "f", None)))
                charts.append((type(chart).__name__, str(chart.title), series))
            payload["sheets"].append({"name": sheet.title, "state": sheet.sheet_state, "dimension": sheet.calculate_dimension(), "cells": cells, "tables": tables, "validations": validations, "charts": charts})
        payload["calculation"] = {
            "mode": formula_book.calculation.calcMode,
            "full_calc_on_load": formula_book.calculation.fullCalcOnLoad,
            "force_full_calc": formula_book.calculation.forceFullCalc,
        }
        encoded = json.dumps(payload, sort_keys=True, separators=(",", ":"), default=str).encode("utf-8")
        return hashlib.sha256(encoded).hexdigest()
    finally:
        formula_book.close()
        data_book.close()


def _run_excel_recalculation(repository_root: Path, candidate: Path, timeout_seconds: int = 150) -> None:
    script = repository_root / "scripts/recalculate_workbook.ps1"
    completed = subprocess.run(
        [
            "powershell.exe",
            "-NoProfile",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(script),
            "-WorkbookPath",
            str(candidate),
            "-TimeoutSeconds",
            str(timeout_seconds),
        ],
        cwd=repository_root,
        capture_output=True,
        text=True,
        timeout=timeout_seconds + 30,
        check=False,
    )
    if completed.returncode != 0:
        details = " ".join(
            line.strip()
            for line in (completed.stderr or completed.stdout).splitlines()
            if line.strip()
        )
        safe_details = details.replace(str(candidate), "<candidate>").replace(
            str(repository_root), "<repository>"
        )
        suffix = f" Detail: {safe_details}" if safe_details else ""
        raise RuntimeError(
            "Desktop Excel recalculation did not complete successfully." + suffix
        )
    process_match = re.search(
        r"DedicatedProcessId=(\d+)", completed.stdout, flags=re.IGNORECASE
    )
    if process_match is None:
        raise RuntimeError(
            "Desktop Excel recalculation did not report its dedicated process identity."
        )
    excel_process_id = int(process_match.group(1))
    exit_deadline = time.monotonic() + 30
    while _windows_process_is_running(excel_process_id):
        if time.monotonic() >= exit_deadline:
            raise RuntimeError(
                "Desktop Excel recalculation completed, but its dedicated process did not exit."
            )
        time.sleep(0.25)


def _windows_process_is_running(process_id: int) -> bool:
    """Return whether a Windows process is still active without terminating it."""

    process_query_limited_information = 0x1000
    still_active = 259
    kernel32 = ctypes.windll.kernel32
    handle = kernel32.OpenProcess(
        process_query_limited_information, False, process_id
    )
    if not handle:
        return False
    try:
        exit_code = ctypes.c_ulong()
        if not kernel32.GetExitCodeProcess(handle, ctypes.byref(exit_code)):
            return False
        return exit_code.value == still_active
    finally:
        kernel32.CloseHandle(handle)


def build_verified_workbook(repository_root: Path, final_path: Path | None = None) -> dict[str, object]:
    """Build, recalculate, verify, and atomically publish the workbook."""

    destination = final_path or repository_root / WORKBOOK_PATH
    with tempfile.TemporaryDirectory(prefix="nike-excel-model-") as temp_dir:
        candidate = Path(temp_dir) / "nike_valuation_model.candidate.xlsx"
        build_result = build_workbook(repository_root, candidate)
        inspect_workbook(candidate, repository_root, require_recalculated=False)
        _run_excel_recalculation(repository_root, candidate)
        normalize_workbook_package(candidate)
        verification = inspect_workbook(candidate, repository_root, require_recalculated=True)
        digest = semantic_digest(candidate)
        destination.parent.mkdir(parents=True, exist_ok=True)
        publish_candidate = destination.with_suffix(".publishing.xlsx")
        shutil.copy2(candidate, publish_candidate)
        os.replace(publish_candidate, destination)
    return {**build_result, "verification": verification, "semantic_digest": digest, "path": destination}


def _parse_args(arguments: Iterable[str] | None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build and verify the Phase 5B Excel valuation model.")
    parser.add_argument("--verify-only", action="store_true", help="Verify the published workbook without rebuilding it.")
    return parser.parse_args(arguments)


def main(arguments: Iterable[str] | None = None) -> int:
    args = _parse_args(arguments)
    root = find_repository_root(Path.cwd())
    workbook_path = root / WORKBOOK_PATH
    if args.verify_only:
        result = inspect_workbook(workbook_path, root, require_recalculated=True)
        print(f"Workbook verified: {len(result['sheets'])} visible sheets; status {result['check_status']}.")
        return 0
    result = build_verified_workbook(root, workbook_path)
    verification = result["verification"]
    print("Phase 5B workbook built, recalculated in desktop Excel, and verified.")
    print(f"Sheets: {len(verification['sheets'])}; formulas: {verification['formula_count']}.")
    print(f"Checks: {verification['pass_count']} pass; {verification['warning_count']} warning; {verification['fail_count']} fail.")
    print(f"Semantic digest: {result['semantic_digest']}.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
