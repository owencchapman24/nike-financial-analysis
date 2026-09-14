from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory

import openpyxl

from nike_financial_analysis import excel_model
from nike_financial_analysis.excel_model import (
    WORKBOOK_PATH,
    build_workbook,
    inspect_workbook,
    normalize_workbook_package,
    semantic_digest,
    verify_all_protected_artifacts,
)
from nike_financial_analysis.valuation import build_valuation_model


ROOT = Path(__file__).resolve().parents[1]
SCENARIO_CHECK_IDS = {
    "terminal_bridge_revenue",
    "terminal_bridge_fcff",
    "terminal_value_formula",
    "enterprise_value_equation",
    "equity_value_equation",
    "basic_share_crosscheck",
}


def _number(workbook: openpyxl.Workbook, sheet: str, cell: str) -> Decimal:
    return Decimal(str(workbook[sheet][cell].value))


def _select_scenario(path: Path, scenario: str) -> None:
    workbook = openpyxl.load_workbook(path, data_only=False, keep_links=True)
    try:
        workbook["Cover"]["D6"] = scenario
        workbook.calculation.calcMode = "auto"
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.save(path)
    finally:
        workbook.close()


def _check_rows(workbook: openpyxl.Workbook) -> dict[str, int]:
    checks = workbook["Checks"]
    return {
        checks[f"A{row}"].value: row
        for row in range(9, checks.max_row + 1)
    }


def test_published_workbook_recalculates_and_reconciles_to_python():
    path = ROOT / WORKBOOK_PATH
    summary = inspect_workbook(path, ROOT, require_recalculated=True)
    model = build_valuation_model(ROOT)
    cached = openpyxl.load_workbook(path, data_only=True, keep_links=True)
    try:
        assert summary["check_status"] == "PASS WITH WARNINGS"
        assert summary["fail_count"] == 0
        assert summary["warning_count"] == 3
        assert _number(cached, "Checks", "D9") == Decimal("2")
        assert abs(_number(cached, "WACC", "E36") - model.wacc_primary.wacc) <= Decimal("1E-9")
        expected = {
            result.scenario: result.value_per_share
            for result in model.scenario_valuations
        }
        for scenario, cell in {"bear": "L26", "base": "M26", "bull": "N26"}.items():
            assert abs(_number(cached, "DCF", cell) - expected[scenario]) <= Decimal("0.01")
        assert abs(_number(cached, "DCF", "E69")) <= Decimal("0.1")
        assert abs(_number(cached, "Sensitivity", "F11") - expected["base"]) <= Decimal("0.01")
    finally:
        cached.close()


def test_published_workbook_has_stable_semantic_digest_and_protected_inputs():
    path = ROOT / WORKBOOK_PATH
    first = semantic_digest(path)
    second = semantic_digest(path)
    assert first == second
    assert len(first) == 64
    assert verify_all_protected_artifacts(ROOT)


def test_every_scenario_passes_after_desktop_excel_recalculation():
    model = build_valuation_model(ROOT)
    expected_values = {
        result.scenario.title(): result.value_per_share
        for result in model.scenario_valuations
    }

    with TemporaryDirectory(prefix="nike-excel-model-") as temp_dir:
        path = Path(temp_dir) / "nike_valuation_model.candidate.xlsx"
        for scenario in ("Bear", "Base", "Bull"):
            build_workbook(ROOT, path)
            _select_scenario(path, scenario)
            excel_model._run_excel_recalculation(ROOT, path)
            normalize_workbook_package(path)
            summary = inspect_workbook(
                path,
                ROOT,
                require_recalculated=True,
                expected_scenario=scenario,
            )
            formulas = openpyxl.load_workbook(
                path, data_only=False, keep_links=True
            )
            cached = openpyxl.load_workbook(path, data_only=True, keep_links=True)
            try:
                rows = _check_rows(formulas)
                assert cached["Cover"]["D6"].value == scenario
                assert cached["Cover"]["D28"].value == "PASS WITH WARNINGS"
                assert cached["Checks"]["H6"].value == "PASS WITH WARNINGS"
                assert summary["pass_count"] == 161
                assert summary["warning_count"] == 3
                assert summary["fail_count"] == 0
                assert Decimal(summary["xnpv_difference"]) == Decimal("0")
                assert abs(
                    _number(cached, "Cover", "D24") - expected_values[scenario]
                ) <= Decimal("0.01")
                assert all(
                    cached["Checks"][f"H{rows[check_id]}"].value == "PASS"
                    for check_id in SCENARIO_CHECK_IDS
                )
            finally:
                formulas.close()
                cached.close()


def test_scenario_expected_value_change_causes_selected_check_failure():
    with TemporaryDirectory(prefix="nike-excel-model-") as temp_dir:
        path = Path(temp_dir) / "nike_valuation_model.candidate.xlsx"
        build_workbook(ROOT, path)
        _select_scenario(path, "Bear")
        workbook = openpyxl.load_workbook(path, data_only=False, keep_links=True)
        try:
            formula = str(workbook["DCF"]["L10"].value)
            workbook["DCF"]["L10"] = f"={formula.removeprefix('=')}+100"
            workbook.save(path)
        finally:
            workbook.close()

        excel_model._run_excel_recalculation(ROOT, path)
        cached = openpyxl.load_workbook(path, data_only=True, keep_links=True)
        formulas = openpyxl.load_workbook(path, data_only=False, keep_links=True)
        try:
            row = _check_rows(formulas)["terminal_bridge_revenue"]
            assert cached["Checks"][f"H{row}"].value == "FAIL"
            assert cached["Checks"]["H6"].value == "FAIL"
            assert cached["Cover"]["D28"].value == "FAIL"
        finally:
            cached.close()
            formulas.close()
