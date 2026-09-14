from pathlib import Path

import conftest as pytest_configuration
import openpyxl
import pytest

from nike_financial_analysis import excel_model
from nike_financial_analysis.excel_model import (
    REQUIRED_DEFINED_NAMES,
    SHEET_ORDER,
    build_workbook,
    inspect_workbook,
    normalize_workbook_package,
)


ROOT = Path(__file__).resolve().parents[1]


def test_formula_workbook_has_approved_structure_and_native_logic(tmp_path):
    path = tmp_path / "candidate.xlsx"
    build_workbook(ROOT, path)
    normalize_workbook_package(path)

    summary = inspect_workbook(path, ROOT, require_recalculated=False)
    workbook = openpyxl.load_workbook(path, data_only=False, keep_links=True)
    try:
        assert tuple(workbook.sheetnames) == SHEET_ORDER
        assert all(workbook[name].sheet_state == "visible" for name in SHEET_ORDER)
        assert REQUIRED_DEFINED_NAMES <= set(workbook.defined_names)
        assert workbook["Cover"]["D6"].value == "Base"
        validation = workbook["Cover"].data_validations.dataValidation[0]
        assert validation.formula1 == "ScenarioList"
        assert "D6" in str(validation.sqref)
        assert workbook["WACC"]["E36"].data_type == "f"
        assert workbook["DCF"]["E59"].data_type == "f"
        assert workbook["Sensitivity"]["F11"].data_type == "f"
        checks = workbook["Checks"]
        check_rows = {
            checks[f"A{row}"].value: row
            for row in range(9, checks.max_row + 1)
        }
        scenario_expected_rows = {
            "terminal_bridge_revenue": 10,
            "terminal_bridge_fcff": 20,
            "terminal_value_formula": 22,
            "enterprise_value_equation": 24,
            "equity_value_equation": 25,
            "basic_share_crosscheck": 27,
        }
        for check_id, dcf_row in scenario_expected_rows.items():
            row = check_rows[check_id]
            assert str(checks[f"C{row}"].value).endswith("selected scenario")
            assert checks[f"E{row}"].value == (
                "=CHOOSE(MATCH(SelectedScenario,ScenarioList,0),"
                f"'DCF'!L{dcf_row},'DCF'!M{dcf_row},'DCF'!N{dcf_row})"
            )
        selected_block = range(
            check_rows["discount_exponent_FY2027"],
            check_rows["basic_share_crosscheck"] + 1,
        )
        assert all(
            not str(checks[f"C{row}"].value).endswith("selected Base")
            for row in selected_block
        )
        formulas = [
            str(cell.value)
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if cell.data_type == "f"
        ]
        assert summary["formula_count"] >= 1_000
        assert any("XNPV(" in formula.upper() for formula in formulas)
        assert not any("INDIRECT(" in formula.upper() for formula in formulas)
        assert not any("OFFSET(" in formula.upper() for formula in formulas)
    finally:
        workbook.close()


def test_sensitivity_and_control_formulas_preserve_live_dependencies(tmp_path):
    path = tmp_path / "candidate.xlsx"
    build_workbook(ROOT, path)
    workbook = openpyxl.load_workbook(path, data_only=False, keep_links=True)
    try:
        sensitivity = workbook["Sensitivity"]
        for column in "DEFGH":
            assert sensitivity[f"{column}18"].value == f"={column}$8"
        for row in range(9, 14):
            for column in "DEFGH":
                formula = str(sensitivity[f"{column}{row}"].value)
                assert "BaseValuePerShare" not in formula
                assert f"$C{row}" in formula
                assert f"{column}$8" in formula
                assert f"{column}29" in formula

        checks = workbook["Checks"]
        rows = {
            checks[f"A{row}"].value: row
            for row in range(9, checks.max_row + 1)
        }
        for row in range(9, 14):
            for column in "DEFGH":
                calculation_row = rows[f"sensitivity_calculation_{column}{row}"]
                assert checks[f"B{calculation_row}"].value == "Mechanical integrity"
                assert checks[f"D{calculation_row}"].value == (
                    f"='Sensitivity'!{column}{row}"
                )
                independent_formula = str(checks[f"E{calculation_row}"].value)
                assert f"'Sensitivity'!$C${row}" in independent_formula
                assert f"'Sensitivity'!${column}$8" in independent_formula
                assert f"'Sensitivity'!{column}{row}" not in independent_formula
                assert f"'Sensitivity'!${column}$29" not in independent_formula
                assert (
                    checks[f"B{rows[f'sensitivity_{row}_{column}']}"].value
                    == "Approved snapshot"
                )

        center_row = rows["sensitivity_center"]
        assert checks[f"B{center_row}"].value == "Conditional / input warning"
        assert checks[f"E{center_row}"].value == "=BaseValuePerShare"
        assert "COORDINATES DIFFER" in str(checks[f"H{center_row}"].value)
        assert "ABS('Sensitivity'!$C$11-CalculatedWACC)<=1E-9" in str(
            checks[f"H{center_row}"].value
        )
        assert "ABS('Sensitivity'!$F$8-TerminalGrowth)<=1E-9" in str(
            checks[f"H{center_row}"].value
        )

        snapshot_row = rows["valuation_input_values_match_approved"]
        assert checks[f"B{snapshot_row}"].value == "Approved snapshot"
        assert "MATCHES APPROVED" in str(checks[f"H{snapshot_row}"].value)
        assert "DIFFERS FROM APPROVED" in str(checks[f"H{snapshot_row}"].value)
        source_row = rows["valuation_source_id_resolution"]
        source_formula = str(checks[f"D{source_row}"].value)
        assert "LEN(TRIM('Sources'!G7:G40))>0" in source_formula
        assert 'SEARCH(";;"' in source_formula
        assert '";"&\'Sources\'!$A$45&";"' in source_formula
        assert ',"VS01","")' not in source_formula
        assert 'SUBSTITUTE(TRIM(\'Sources\'!G7:G40)," ","")' not in source_formula
        syntax_row = rows["valuation_source_id_syntax"]
        assert checks[f"B{syntax_row}"].value == "Package/build control"
        assert 'MID(\'Sources\'!A45:A55,3,1)' in str(
            checks[f"D{syntax_row}"].value
        )

        fcff_row = rows["fcff_identity_base_FY2027"]
        assert checks[f"B{fcff_row}"].value == "Mechanical integrity"
        assert "'Scenarios'!E52-(" in str(checks[f"D{fcff_row}"].value)
        assert checks[f"E{fcff_row}"].value == 0
        assert workbook.defined_names["MechanicalIntegrityStatus"].attr_text == (
            "'Checks'!$A$6"
        )
        assert workbook.defined_names["ApprovedSnapshotStatus"].attr_text == (
            "'Checks'!$C$6"
        )
        assert workbook.defined_names["PackageBuildControlStatus"].attr_text == (
            "'Checks'!$G$6"
        )
        assert workbook.defined_names["ModelStatus"].attr_text == "'Checks'!$H$6"
        assert checks["G5"].value == "Package/build controls"
        assert checks["H5"].value == "Canonical publication gate"
        assert workbook["Cover"]["D28"].value == "=ModelStatus"
        assert checks.column_dimensions["A"].width >= 48
        assert checks.column_dimensions["G"].width >= 26
        assert checks.column_dimensions["H"].width >= 58
    finally:
        workbook.close()


def test_pre_excel_recalculation_sentinel_is_deliberately_invalid(tmp_path):
    path = tmp_path / "candidate.xlsx"
    build_workbook(ROOT, path)
    cached = openpyxl.load_workbook(path, data_only=True)
    try:
        assert cached["Checks"]["D9"].value == -999
        assert cached["Cover"]["D8"].value.date().isoformat() == "2026-09-04"
        assert cached["Sources"]["N13"].value.date().isoformat() == "2026-09-04"
        assert cached["DCF"]["D12"].value.date().isoformat() == "2026-09-04"
    finally:
        cached.close()


def test_fail_closed_build_does_not_publish_unverified_workbook(tmp_path, monkeypatch):
    destination = tmp_path / "model" / "nike_valuation_model.xlsx"

    def fail_recalculation(*_args, **_kwargs):
        raise RuntimeError("controlled recalculation failure")

    monkeypatch.setattr(excel_model, "_run_excel_recalculation", fail_recalculation)
    with pytest.raises(RuntimeError, match="controlled recalculation failure"):
        excel_model.build_verified_workbook(ROOT, destination)
    assert not destination.exists()
    assert not destination.with_suffix(".publishing.xlsx").exists()


def test_recalculation_script_and_excel_opt_in_are_bounded(
    pytestconfig, monkeypatch
):
    script = (ROOT / "scripts/recalculate_workbook.ps1").read_text(encoding="utf-8")
    assert "CalculateFullRebuild" in script
    assert "TimeoutSeconds" in script
    assert "DedicatedProcessId=" in script
    assert "DisplayAlerts = $false" in script
    assert "EnableEvents = $false" in script
    assert "AskToUpdateLinks = $false" in script
    assert "CircularReference" in script
    assert "Stop-Process" not in script
    assert "taskkill" not in script.lower()

    marker_entries = [
        marker
        for marker in pytestconfig.getini("markers")
        if marker.partition(":")[0].strip() == "excel_integration"
    ]
    assert len(marker_entries) == 1

    class OptedInConfig:
        @staticmethod
        def getoption(name):
            assert name == "--run-excel-integration"
            return True

    class IntegrationItem:
        @staticmethod
        def get_closest_marker(name):
            return object() if name == "excel_integration" else None

    monkeypatch.setattr(
        pytest_configuration,
        "_excel_prerequisite_errors",
        lambda: ["controlled missing prerequisite"],
    )
    with pytest.raises(
        pytest.UsageError,
        match="controlled missing prerequisite.*omit --run-excel-integration",
    ):
        pytest_configuration.pytest_collection_modifyitems(
            OptedInConfig(), [IntegrationItem()]
        )
